# -*- coding: utf-8 -*-
"""Extract all tables + captions + notes from 手稿1.docx into 手稿1_表格.docx."""
import re
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.table import Table
from docx.text.paragraph import Paragraph

SRC = r"E:/workbuddy/Front  in  Endocrinology  投稿/手稿1.docx"
OUT = r"E:/workbuddy/Front  in  Endocrinology  投稿/手稿1_表格.docx"

src = Document(SRC)

def is_caption(p):
    t = p.text.strip()
    return bool(re.match(r"^(Supplementary )?Table \S+\.", t))

sections, cur, state = [], None, None
for child in src.element.body.iterchildren():
    if child.tag == qn("w:p"):
        p = Paragraph(child, src)
        t = p.text.strip()
        if is_caption(p):
            cur = {"caption": t, "elems": []}
            sections.append(cur); state = "cap"; continue
        if cur is None or not t:
            continue
        if state == "done":
            continue
        if re.match(r"^\([abc]\)\s", t):
            cur["elems"].append(("sub", t)); state = "sub"; continue
        if state in ("cap", "tbl", "sub") and False:
            pass
        if state in ("cap",):
            # first content paragraph after caption that is not a table/sub -> note
            cur["elems"].append(("note", t.strip("*").strip())); state = "done"; continue
        cur["elems"].append(("note", t.strip("*").strip())); state = "done"
    elif child.tag == qn("w:tbl") and cur is not None and state != "done":
        cur["elems"].append(("tbl", Table(child, src))); state = "tbl"

print("sections found:", len(sections))
for s in sections:
    kinds = [(k, len(v.rows) if k == "tbl" else v[:30]) for k, v in s["elems"]]
    print(" -", s["caption"][:44], kinds)

# ---------------- rebuild ----------------
out = Document()
stn = out.styles["Normal"]
stn.font.name = "Times New Roman"; stn.font.size = Pt(10.5)
stn.element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
stn.paragraph_format.space_after = Pt(4)

def para(text, size=10.5, bold=False, align=None):
    p = out.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    r.font.size = Pt(size); r.font.bold = bold
    r.font.name = "Times New Roman"
    r.element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    return p

def set_three_line(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge, sz, val in (("top", 12, "single"), ("bottom", 12, "single"),
                          ("left", 0, "none"), ("right", 0, "none"),
                          ("insideH", 0, "none"), ("insideV", 0, "none")):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), val); el.set(qn("w:sz"), str(sz))
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tblPr.append(borders)
    for cell in table.rows[0].cells:
        tcPr = cell._tc.get_or_add_tcPr()
        tb = OxmlElement("w:tcBorders")
        bot = OxmlElement("w:bottom")
        bot.set(qn("w:val"), "single"); bot.set(qn("w:sz"), "6")
        bot.set(qn("w:color"), "000000")
        tb.append(bot); tcPr.append(tb)

def copy_table(src_tbl, fontsize=8.5):
    nrow = len(src_tbl.rows); ncol = len(src_tbl.columns)
    t = out.add_table(rows=nrow, cols=ncol)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, srow in enumerate(src_tbl.rows):
        for ci, scell in enumerate(srow.cells):
            cell = t.cell(ri, ci)
            cp = cell.paragraphs[0]
            cp.paragraph_format.space_after = Pt(1)
            for sp in scell.paragraphs:
                for run in sp.runs:
                    r = cp.add_run(run.text)
                    r.font.size = Pt(fontsize)
                    r.font.bold = (ri == 0) or run.font.bold
                    r.font.name = "Times New Roman"
                    r.element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    set_three_line(t)
    return t

para("手稿1 —— Tables and table notes", 13, bold=True,
     align=WD_ALIGN_PARAGRAPH.CENTER)
para("Table 1-5 主体表格与 Supplementary Tables S1–S6；表题在上方，表注在下方（三线表）。", 9)

for s in sections:
    para(s["caption"], 9.5, bold=True)
    for kind, val in s["elems"]:
        if kind == "sub":
            para(val, 8.5, bold=True)
        elif kind == "tbl":
            copy_table(val)
        elif kind == "note":
            para(val, 8.5)
out.save(OUT)
print("saved:", OUT)
