# -*- coding: utf-8 -*-
"""
phase9c_text_mode.py
--------------------
T 模式（文本模式）盲法二次复核：阅片包生成 + 裁决表生成 + 预设统计 + 自检。

角色（三位阅片者）
  R1  原始报告的分级者（既有，其结论以报告文本形式存在）
  R2  独立盲法阅片者（只读报告文本）
  R3  独立盲法阅片者（与 R2 互不知晓，同一份盲号与顺序）
  R4  裁决者（视网膜专科；只看分歧记录，裁定分级）
  —— R2/R3 之间用于"阅片者间一致性"；R1–R2、R1–R3 用于"我们的推导是否忠实"；
     R4 的裁定结果用于重算患病率与层 2 判别力。

子命令
  package                   生成两份阅片包（R2 用、R3 用）+ 练习答案
  adjudicate --r2 --r3      比对两份回收件，生成只含分歧记录的裁决表（发给 R4）
  analyse --r2 --r3 [--adj] 计算全部预设统计量，输出 json 与分歧矩阵
  selftest                  端到端自检（模拟 R2/R3/R4）并打印关键统计

T 模式的定义（写进 Methods 时必须准确）
  阅片者**只读原始报告文本**（不看眼底照片），独立给出 ICDR 五级分级、any-DR
  与"是否可分级"。它检验的是"我们**从报告推导分级**这一步是否忠实"，
  **不是**"报告本身是否正确"——后者需要影像模式（I 模式）。
"""
import sys, os, json, argparse
sys.stdout.reconfigure(encoding='utf-8')
import numpy as np
import pandas as pd
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

SRC = r"E:/workbuddy/2026-09-04-18-19-16/phase2_dev"
REPO = r"E:/workbuddy/lipid-DR-four-layer-falsification"
REPO_OT = os.path.join(REPO, "output_tables")
OUTDIR = r"E:/workbuddy/Front  in  Endocrinology  投稿"
SAMPLING = os.path.join(REPO_OT, "layer2_second_grader_sampling.csv")
PKG_R2 = os.path.join(OUTDIR, "_T模式_阅片包_R2_20260929.xlsx")
PKG_R3 = os.path.join(OUTDIR, "_T模式_阅片包_R3_20260929.xlsx")
ADJ_OUT = os.path.join(OUTDIR, "_T模式_裁决表_R4_20260929.xlsx")
SEED = 20260929
NBOOT = 2000
SHEET2 = '2_GradingSheet_正式阅片'
SHEET_ADJ = '2_Adjudication'

GRADE_W = np.array([[(i - j) ** 2 / 16 for j in range(5)] for i in range(5)])


# ---------------------------------------------------------------- 数据
def load_base():
    aud = pd.read_csv(f"{SRC}/severity_rebuild_audit.csv", low_memory=False)
    dev = pd.read_csv(f"{SRC}/dev_set_v2.csv", low_memory=False)
    assert len(aud) == len(dev)
    return pd.DataFrame({
        'pid': aud['pid'].astype(str), 'year': aud['year'],
        'Ltxt': aud['Ltxt'].fillna(''), 'Rtxt': aud['Rtxt'].fillna(''),
        'gradeable': dev['gradeable'].values, 'any_dr': dev['any_dr'].values,
        'severity': dev['severity'].values,
    }), dev


def master_key():
    """盲号 ↔ 分层/原分级 对照；顺序与阅片包一致（同一种子）。"""
    samp = pd.read_csv(SAMPLING)
    samp['pid'] = samp.pid.astype(str)
    k = samp.sample(frac=1, random_state=SEED).reset_index(drop=True)
    k['blind_id'] = ['SG-%03d' % (i + 1) for i in range(len(k))]
    return k


# ---------------------------------------------------------------- 统计工具
def cohen_kappa(a, b, weights=None, categories=None):
    a, b = np.asarray(a), np.asarray(b)
    cats = categories if categories is not None else sorted(set(a) | set(b))
    idx = {c: i for i, c in enumerate(cats)}
    O = np.zeros((len(cats), len(cats)))
    for x, y in zip(a, b):
        O[idx[x], idx[y]] += 1
    O /= O.sum()
    E = np.outer(O.sum(1), O.sum(0))
    W = np.ones(O.shape) - np.eye(O.shape[0]) if weights is None else np.asarray(weights)
    den = (W * E).sum()
    return np.nan if den == 0 else 1 - (W * O).sum() / den


def boot_ci(func, n, B=NBOOT, seed=SEED):
    rng = np.random.default_rng(seed)
    vals = []
    for _ in range(B):
        i = rng.integers(0, n, n)
        v = func(i)
        if v is not None and not (isinstance(v, float) and np.isnan(v)):
            vals.append(v)
    v = np.array(vals, float)
    return (np.nan, np.nan) if len(v) < 10 else tuple(np.percentile(v, [2.5, 97.5]))


def pabak(a, b, cats):
    return (len(cats) * np.mean(np.asarray(a) == np.asarray(b)) - 1) / (len(cats) - 1)


def auc_score(y, s):
    y = np.asarray(y, float); s = np.asarray(s, float)
    r = pd.Series(s).rank().values
    n1, n0 = y.sum(), len(y) - y.sum()
    return np.nan if n0 == 0 or n1 == 0 else (r[y == 1].sum() - n1 * (n1 + 1) / 2) / (n0 * n1)


def pair_stats(P_grad, R_grad, P_any, R_any, P5, R5, tag, out):
    """一对阅片者之间的全部一致性指标（P_* 为左侧，R_* 为右侧）。"""
    both = (P_grad == 1) & (R_grad == 1)
    o = {}
    o['gradable_agree'] = int((P_grad == R_grad).sum())
    o['gradable_kappa'] = round(float(cohen_kappa(P_grad, R_grad, categories=[0, 1])), 4)
    o['gradable_pabak'] = round(float(pabak(P_grad, R_grad, [0, 1])), 4)
    o['right_gradable_left_excluded'] = int(((R_grad == 1) & (P_grad == 0)).sum())
    o['left_gradable_right_excluded'] = int(((R_grad == 0) & (P_grad == 1)).sum())

    pa = np.where(P_grad == 0, 2, np.where(P_any > 0, 1, 0))
    ra = np.where(R_grad == 0, 2, np.where(R_any > 0, 1, 0))
    o['anydr3_agree'] = int((pa == ra).sum())
    o['anydr3_kappa'] = round(float(cohen_kappa(pa, ra, categories=[0, 1, 2])), 4)
    o['anydr3_pabak'] = round(float(pabak(pa, ra, [0, 1, 2])), 4)

    pb, rb = P_any[both], R_any[both]
    o['anydr_n_both_gradable'] = int(both.sum())
    o['anydr_agree'] = int((pb == rb).sum())
    o['anydr_kappa'] = round(float(cohen_kappa(pb, rb, categories=[0, 1])), 4)
    o['anydr_pabak'] = round(float(pabak(pb, rb, [0, 1])), 4)
    o['anydr_kappa_CI'] = [round(v, 4) for v in boot_ci(
        lambda i: cohen_kappa(pb[i], rb[i], categories=[0, 1]), max(len(pb), 1))]
    o['flips_to_case'] = int(((pb == 0) & (rb == 1)).sum())
    o['flips_to_noncase'] = int(((pb == 1) & (rb == 0)).sum())

    p5, r5 = P5[both], R5[both]
    o['five_level_n'] = int(both.sum())
    o['five_level_agree'] = int((p5 == r5).sum())
    o['five_level_wkappa'] = round(float(cohen_kappa(p5, r5, weights=GRADE_W,
                                                    categories=[0., 1., 2., 3., 4.])), 4)
    o['five_level_pabak'] = round(float(pabak(p5, r5, [0., 1., 2., 3., 4.])), 4)
    o['five_level_wkappa_CI'] = [round(v, 4) for v in boot_ci(
        lambda i: cohen_kappa(p5[i], r5[i], weights=GRADE_W, categories=[0., 1., 2., 3., 4.]),
        max(len(p5), 1))]
    d = np.abs(p5 - r5)
    o['five_level_1grade_diff'] = int((d == 1).sum())
    o['five_level_ge2grade_diff'] = int((d >= 2).sum())
    out[tag] = o
    return o


# ---------------------------------------------------------------- package
INSTR_COMMON = [
    ("一、这次请您做什么", ""),
    ("", "请依据每份报告文本，独立给出该次筛查记录的眼底分级。本队列的报告文本由眼科医师在标准眼底照片上判读后书写；"
         "本次**不提供照片**，只提供文字，因此您判读的是**报告文本所表达的分级**，而不是重新看片。"),
    ("", "每条记录请填写 4 项：①左眼分级；②右眼分级；③该记录的总体分级（取较重眼）；④该记录是否为 any-DR（1 = 有，0 = 无）。"
         "若这条报告的文字不足以支持一个明确分级，请在\"可分级\"列填 0，**此时分级列与 any-DR 列请留空**，并在备注说明原因。"),
    ("二、分级标准（ICDR）", ""),
    ("", "0 = 无明显糖尿病视网膜改变；1 = 轻度非增殖期（mild NPDR）；2 = 中度非增殖期（moderate NPDR）；"
         "3 = 重度非增殖期（severe NPDR）；4 = 增殖期（PDR）。记录级分级取**较重眼**（worst-eye rule）。"),
    ("三、本队列的常用报告用语（见第 3 张表）", ""),
    ("", "本队列报告中\"糖尿病视网膜改变\"与\"糖尿病视网膜病变\"两种写法混用，含义相同；"
         "非增殖期的严重度写作\"轻度／中度／重度非增殖性\"。请以第 3 张表的对应关系为准。"),
    ("四、四处需要您按临床判断处理的地方", ""),
    ("", "1. 报告只写\"可疑糖尿病视网膜病变\"而未给出分级：请按您的临床判断填写，并在备注写明理由。"
         "**本次不预设规则**——我们正是想知道临床判读者会怎么处理。"),
    ("", "2. 报告同时出现\"未见明显糖尿病视网膜改变\"与\"可见微血管瘤、可见出血\"等阳性描述：请以阳性描述为准判断，并备注。"),
    ("", "3. 报告仅给出自动量化数值（出血数量／出血总面积／渗出数量等）而无任何分级或定性描述：请填\"不可分级\"（0），"
         "因为这类输出无法映射到 ICDR 分级。"),
    ("", "4. 少数记录只有一只眼有报告文本（本样本共 6 条）：请按有文本的那只眼给出分级与 any-DR，并备注另一眼无文本。"),
    ("五、阅读方式（盲法要求）", ""),
    ("", "① 请只打开第 2 张表；内部对照表**不在本文件内**，请勿索取。"
         "② 记录顺序已随机打乱，请**按行顺序逐条判读，不要回看**已判过的记录。"
         "③ 建议一次完成。④ 每条独立判断，不参考前后记录。"),
    ("", "⑤ 本批次由多位阅片者各自独立完成，**请勿与任何人讨论任何一条记录**——这是本次结果能否成立的前提。"),
    ("六、先做练习（第 1 张表）", ""),
    ("", "第 1 张表是 20 例练习，**已给出参考答案**（最右几列），供您对齐标准。"
         "练习题取自同一队列但**不在正式阅片样本中**。建议先做完练习、核对答案，再开始正式阅片。"),
    ("七、回传", ""),
    ("", "请在原文件上填写第 1、2 张表的空列后回传该 xlsx。请勿改动盲号列、文本列与行顺序。"
         "填完后将计算五级一致性与 any-DR 一致性（二次加权 kappa 与未加权 kappa），以及**阅片者之间**的一致性；"
         "分歧由第三位视网膜专科医师裁定，裁定前后的结果都会报告。"),
    ("English method note (for the record)", ""),
    ("", "Two readers independently assigned the ICDR severity grade (0-4) and the any-DR endpoint for each record "
         "from the source report text alone, without access to the fundus photographs, to the primary grade, to the "
         "primary binary outcome, to the lipid values or to any interim result, and without knowledge of each other's "
         "grades. The worst eye defined the record-level grade. Records were presented in random order under a blind "
         "identifier. Twenty practice records drawn from the same cohort but outside the analysis sample were provided, "
         "with their key, for calibration and were not analysed. Readers were given the ICDR scale and the cohort's "
         "standard report vocabulary, but no decision rule for reports that state suspected retinopathy without a "
         "grade, so that the handling of that wording could be observed rather than prescribed. Where the text did not "
         "support an explicit grade, the record was recorded as not gradable. Disagreements between the two readers are "
         "resolved by a third reader, a retinal specialist independent of both."),
]


def build_package():
    D, dev = load_base()
    samp = pd.read_csv(SAMPLING)
    in_sample = set(samp.pid.astype(str))
    g = D.gradeable.eq(1)

    # 练习集：同队列、不在样本内（G1/G3/G4 已 100% 入样，故只能从 G0/G2 抽）
    pool = D[g & ~D.pid.isin(in_sample)]
    practice = pd.concat([pool[pool.severity.eq(0)].sample(10, random_state=1),
                          pool[pool.severity.eq(2)].sample(10, random_state=2)]) \
        .sample(frac=1, random_state=SEED).reset_index(drop=True)
    practice['practice_id'] = ['P-%02d' % (i + 1) for i in range(len(practice))]
    assert not (set(practice.pid) & in_sample), "练习集与正式样本重叠"

    key = master_key()
    sheet1 = pd.DataFrame({
        'Practice ID': practice.practice_id, 'Year': practice.year,
        'Left eye report text': practice.Ltxt, 'Right eye report text': practice.Rtxt,
        'Your grade LEFT (0-4)': '', 'Your grade RIGHT (0-4)': '',
        'Your worst-eye grade (0-4)': '', 'Your any-DR (1/0)': '',
        'Your gradable? (1/0)': '', 'Comment': '',
        'KEY: worst-eye grade (0-4)': practice.severity,
        'KEY: any-DR (1/0)': practice.any_dr,
        'KEY: gradable (1/0)': practice.gradeable,
    })
    sheet2 = pd.DataFrame({
        'Blind ID': key.blind_id, 'Screening year': key.year,
        'Left eye report text': key.Ltxt, 'Right eye report text': key.Rtxt,
        'Grade LEFT (0-4)': '', 'Grade RIGHT (0-4)': '', 'Worst-eye grade (0-4)': '',
        'Any DR (1=yes, 0=no)': '', 'Gradable? (1=yes, 0=no)': '',
        'Confidence (1 high / 2 medium / 3 low)': '', 'Comment': '',
    })
    sheet3 = pd.DataFrame({
        'Report wording (as it appears in this cohort)': [
            '未见明显糖尿病视网膜改变 / 未见明显糖尿病视网膜病变',
            '未见微血管瘤,未见出血,未见硬性渗出,……,未见明显糖尿病视网膜病变',
            '轻度非增殖性糖尿病视网膜改变／病变',
            '中度非增殖性糖尿病视网膜改变／病变',
            '重度非增殖性糖尿病视网膜改变／病变',
            '增殖性糖尿病视网膜病变／改变（或列出新生血管、玻璃体积血等）',
            '出血数量：N；出血总面积：N；渗出数量：N；……（仅自动量化）',
            '屈光介质混浊，眼底看不清',
            '可见糖尿病视网膜病变（提及但未给分级）',
            '可疑糖尿病视网膜病变（未给分级）',
            '有其他眼底异常'],
        'Grade to assign (0-4)': ['0', '0', '1', '2', '3', '4', 'not gradable', 'not gradable',
                                  'not gradable (no grade stated)',
                                  'your clinical judgement - see instruction 1',
                                  'no grade from this phrase alone'],
        'Notes': ['explicit absence of retinopathy', 'explicit absence, itemised negative descriptors',
                  'mild non-proliferative', 'moderate non-proliferative', 'severe non-proliferative',
                  'proliferative', 'automated quantitative output only; maps to no ICDR grade',
                  'media opacity; not gradable', 'retinopathy mentioned but no severity grade stated',
                  'suspected; not equivalent to absence - record your own judgement',
                  'this phrase alone carries no grade'],
    })

    for path, role in [(PKG_R2, '阅片者 2（R2）'), (PKG_R3, '阅片者 3（R3）')]:
        instr = [("本文件的使用者", "本文件是为**%s**准备的。另有一位阅片者持有一份内容相同但独立填写的文件；"
                                    "两位**互不知晓对方的判读结果**，请勿相互讨论。" % role)] + INSTR_COMMON
        s0 = pd.DataFrame(instr, columns=['请阅读：T 模式盲法复核说明', '内容'])
        with pd.ExcelWriter(path, engine='openpyxl') as xw:
            s0.to_excel(xw, sheet_name='0_Instructions_请先阅读', index=False)
            sheet1.to_excel(xw, sheet_name='1_Practice_20_含答案', index=False)
            sheet2.to_excel(xw, sheet_name=SHEET2, index=False)
            sheet3.to_excel(xw, sheet_name='3_PhraseGuide_用语对照', index=False)
        _format(path)
        print("已写出阅片包（%s）: %s" % (role, os.path.basename(path)))

    practice.to_csv(os.path.join(REPO_OT, 'layer2_text_mode_practice_key.csv'),
                    index=False, encoding='utf-8-sig')
    print("正式阅片表 %d 例；练习 %d 例（不入统计，与正式样本无重叠）" % (len(sheet2), len(practice)))


def _format(path):
    wb = openpyxl.load_workbook(path)
    for ws in wb.worksheets:
        ws.freeze_panes = 'A2'
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill('solid', fgColor='DDEBF7')
            c.alignment = Alignment(vertical='center', wrap_text=True)
        for j in range(1, ws.max_column + 1):
            mx = max((len(str(c.value or '')) for c in ws[get_column_letter(j)]), default=10)
            ws.column_dimensions[get_column_letter(j)].width = min(max(mx + 2, 10), 60)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, str) and len(c.value) > 45:
                    c.alignment = Alignment(wrap_text=True, vertical='top')
    wb.save(path)


# ---------------------------------------------------------------- 读取回收件


def load_returned(path, sheet=SHEET2, mode='reader'):
    """读回收件。mode='reader' 对应阅片表列名；mode='adj' 对应裁决表列名。"""
    COLS = (dict(grade='Worst-eye grade (0-4)', any='Any DR (1=yes, 0=no)',
                 grad='Gradable? (1=yes, 0=no)') if mode == 'reader' else
            dict(grade='ADJUDICATED grade (0-4)', any='ADJUDICATED any-DR (1/0)',
                 grad='ADJUDICATED gradable (1/0)'))
    x = pd.read_excel(path, sheet_name=sheet)
    need = ['Blind ID'] + list(COLS.values())
    miss = [c for c in need if c not in x.columns]
    assert not miss, "%s 缺少列：%s" % (os.path.basename(path), miss)
    g = pd.to_numeric(x[COLS['grad']], errors='coerce')
    if g.isna().any():
        b = x[g.isna()]
        raise SystemExit("%s：有 %d 条未填'是否可分级'，示例 %s"
                         % (os.path.basename(path), len(b), list(b['Blind ID'][:5])))
    # 判为不可分级者允许分级留空（本就不应给出分级）——一刀切报错会误判
    bad = x[(g > 0) & x[[COLS['grade'], COLS['any']]].isna().any(axis=1)]
    if len(bad):
        raise SystemExit("%s：有 %d 条判为可分级但未给分级/any-DR，示例 %s"
                         % (os.path.basename(path), len(bad), list(bad['Blind ID'][:5])))
    return pd.DataFrame({
        'blind_id': x['Blind ID'],
        'grade': pd.to_numeric(x[COLS['grade']], errors='coerce'),
        'any': pd.to_numeric(x[COLS['any']], errors='coerce'),
        'grad': g.astype(int),
    })


def to_arrays(reader_df, key):
    """按主流程盲号顺序对齐，返回 (gradable, any, grade5)。"""
    m = key[['blind_id']].merge(reader_df, on='blind_id', how='left')
    assert m['grad'].notna().all(), "有盲号未回收：%s" % list(m.loc[m['grad'].isna(), 'blind_id'][:5])
    # 注意：列名 'any' 与 DataFrame.any 方法同名，必须用 m['any'] 取列
    return (m['grad'].astype(int).values,
            m['any'].fillna(0).astype(float).values,
            m['grade'].fillna(-1).astype(float).values)


# ---------------------------------------------------------------- 裁决表
def build_adjudication(f_r2, f_r3):
    key = master_key()
    a2 = to_arrays(load_returned(f_r2), key)
    a3 = to_arrays(load_returned(f_r3), key)
    k = key.copy()
    k['g2'], k['y2'], k['s2'] = a2
    k['g3'], k['y3'], k['s3'] = a3
    disc = (k.g2 != k.g3) | (k.s2 != k.s3)
    d = k[disc].copy()
    sheet = pd.DataFrame({
        'Blind ID': d.blind_id,
        'Left eye report text': d.Ltxt,
        'Right eye report text': d.Rtxt,
        'Reader 2 grade (0-4)': ['' if v < 0 else int(v) for v in d.s2],
        'Reader 2 any-DR': ['' if g == 0 else int(y) for g, y in zip(d.g2, d.y2)],
        'Reader 2 gradable': d.g2,
        'Reader 3 grade (0-4)': ['' if v < 0 else int(v) for v in d.s3],
        'Reader 3 any-DR': ['' if g == 0 else int(y) for g, y in zip(d.g3, d.y3)],
        'Reader 3 gradable': d.g3,
        'ADJUDICATED grade (0-4)': '',
        'ADJUDICATED any-DR (1/0)': '',
        'ADJUDICATED gradable (1/0)': '',
        'Reason for your decision': '',
    })
    note = pd.DataFrame({'Instruction': [
        'You are the adjudicator (R4). You receive only the records on which the two independent readers disagreed.',
        'Please judge each record from the report text yourself and record your own grade; do not simply pick one reader.',
        'ICDR: 0 no apparent retinopathy; 1 mild NPDR; 2 moderate NPDR; 3 severe NPDR; 4 proliferative DR. Record-level grade = worst eye.',
        'If the text does not support an explicit grade, set ADJUDICATED gradable = 0 and leave the grade blank.',
        'The two readers were masked to each other and to the primary grade; you are shown both because your task is to resolve, not to re-read blindly.',
        'Your decisions will be used to recompute prevalence and the layer-2 discrimination, and are reported alongside the pre-adjudication statistics.',
        '中文说明：您是裁决者（R4），只收到两位独立阅片者判读不一致的记录。请自行判读并给出您自己的分级，不要只是二选一。'
        '若文本不足以支持明确分级，请把"ADJUDICATED gradable"填 0 并留空分级列。',
    ]})
    with pd.ExcelWriter(ADJ_OUT, engine='openpyxl') as xw:
        note.to_excel(xw, sheet_name='0_Instructions_裁决说明', index=False)
        sheet.to_excel(xw, sheet_name=SHEET_ADJ, index=False)
    _format(ADJ_OUT)
    print("分歧记录 %d / %d（%.1f%%）→ 裁决表: %s"
          % (len(d), len(k), 100 * len(d) / len(k), os.path.basename(ADJ_OUT)))
    d[['blind_id', 's2', 's3', 'y2', 'y3', 'g2', 'g3']].to_csv(
        os.path.join(REPO_OT, 'layer2_text_mode_discordances.csv'), index=False, encoding='utf-8-sig')
    return ADJ_OUT


# ---------------------------------------------------------------- analyse
def analyse(f_r2, f_r3, f_adj=None, quiet=False):
    D, dev = load_base()
    key = master_key()
    P_grad = key.gradeable.astype(int).values
    P_any = key.any_dr.fillna(0).astype(float).values
    P5 = key.severity.fillna(-1).astype(float).values
    a2 = to_arrays(load_returned(f_r2), key)
    a3 = to_arrays(load_returned(f_r3), key)

    res = {'n_sampled': int(len(key)), 'n_independent_readers': 2, 'n_adjudicators': 1}
    pair_stats(P_grad, a2[0], P_any, a2[1], P5, a2[2], 'primary_vs_R2', res)
    pair_stats(P_grad, a3[0], P_any, a3[1], P5, a3[2], 'primary_vs_R3', res)
    pair_stats(a2[0], a3[0], a2[1], a3[1], a2[2], a3[2], 'R2_vs_R3_inter_reader', res)

    # 共识：两位分级一致则取之，否则待裁决
    same = (a2[0] == a3[0]) & (np.where(a2[0] == 0, -9.0, a2[2]) == np.where(a3[0] == 0, -9.0, a3[2]))
    n_unres = int((~same).sum())
    C_grad = a2[0].copy(); C_any = a2[1].copy(); C5 = a2[2].copy()
    adj_done = bool(f_adj and os.path.exists(f_adj))
    if adj_done:
        am = load_returned(f_adj, sheet=SHEET_ADJ, mode='adj')
        disc_ids = set(key.blind_id.values[~same])
        got = set(am.blind_id)
        if disc_ids - got:
            raise SystemExit("裁决表缺少 %d 条分歧记录，示例 %s"
                             % (len(disc_ids - got), sorted(disc_ids - got)[:5]))
        if got - disc_ids:
            raise SystemExit("裁决表多出 %d 条非分歧记录，示例 %s"
                             % (len(got - disc_ids), sorted(got - disc_ids)[:5]))
        assert len(am) == n_unres, "裁决表行数 %d 与分歧数 %d 不一致" % (len(am), n_unres)
        pos = {b: i for i, b in enumerate(key.blind_id)}
        for _, r in am.iterrows():
            i = pos[r.blind_id]
            C_grad[i] = int(r.grad)
            C_any[i] = 0.0 if r.grad == 0 else float(r['any'])
            C5[i] = -1.0 if r.grad == 0 else float(r.grade)
    res['n_discordant_requiring_adjudication'] = n_unres
    res['consensus_source'] = 'adjudicated (R4)' if adj_done else 'unresolved (R4 not yet returned)'
    res['n_records_with_consensus'] = int(same.sum()) + (n_unres if adj_done else 0)

    # 患病率：以共识分级做后分层加权，复现 24.05%
    N_pop = {0.: 1497, 1.: 49, 2.: 386, 3.: 15, 4.: 24}
    gr = C5[C_grad == 1]
    n_g = {g_: int((gr == g_).sum()) for g_ in N_pop}
    w = np.array([N_pop[g_] / n_g[g_] if n_g.get(g_, 0) > 0 else 0.0 for g_ in gr])
    y = (gr > 0).astype(float)
    res['prevalence_primary_pct'] = 24.05
    if len(y):
        res['prevalence_sample_unweighted_pct'] = round(float(100 * y.mean()), 2)
        res['prevalence_poststratified_pct'] = round(float(100 * np.sum(w * y) / np.sum(w)), 2) if w.sum() else None
        rng = np.random.default_rng(SEED)
        bs = []
        for _ in range(500):
            i = rng.integers(0, len(gr), len(gr))
            gi, wi = gr[i], w[i]
            if wi.sum() == 0:
                continue
            bs.append(float(np.sum(wi * (gi > 0)) / np.sum(wi)))
        if bs:
            lo, hi = np.percentile(bs, [2.5, 97.5])
            res['prevalence_CI_pct'] = [round(100 * lo, 2), round(100 * hi, 2)]

    # 层 2 判别力：用共识分级替换被抽中记录
    import statsmodels.api as sm
    cc = dev[dev.gradeable.eq(1)][['id', 'age', 'male', 'tg', 'tc', 'hdl_c', 'ldl_c',
                                   'apo_a1', 'apo_b', 'any_dr']].dropna().reset_index(drop=True)
    cc['id'] = cc.id.astype(str)
    cc['log2tg'] = np.log2(cc.tg)
    X = cc[['age', 'male', 'log2tg', 'tc', 'hdl_c', 'ldl_c', 'apo_a1', 'apo_b']]

    def fit_auc(yv):
        try:
            m = sm.Logit(yv, sm.add_constant(X, has_constant='add')).fit(disp=0, maxiter=80)
        except Exception:
            return np.nan
        return auc_score(yv, m.predict(sm.add_constant(X, has_constant='add')).values)

    res['auc_primary'] = round(float(fit_auc(cc.any_dr.astype(float))), 4)
    repl_pid = {}
    for i, b in enumerate(key.blind_id):
        if C_grad[i] == 1:
            repl_pid[key.pid.values[i]] = 0.0 if C5[i] <= 0 else 1.0
    yv = pd.Series([repl_pid.get(i, v) for i, v in zip(cc.id, cc.any_dr.astype(float))])
    res['auc_with_consensus'] = round(float(fit_auc(yv)), 4)
    res['auc_change'] = round(res['auc_with_consensus'] - res['auc_primary'], 4)
    res['conclusion_threshold_auc'] = 0.70
    res['conclusion_would_change'] = bool(res['auc_with_consensus'] >= 0.70)

    # 最坏情形：两位阅片者同时与主流程相反者，按相反方向重分类
    both_g = (a2[0] == 1) & (a3[0] == 1)
    to_case = both_g & (P_any == 0) & (a2[1] == 1) & (a3[1] == 1)
    to_non = both_g & (P_any == 1) & (a2[1] == 0) & (a3[1] == 0)
    wc = set(key.pid.values[to_case]); wn = set(key.pid.values[to_non])
    worst = pd.Series([1.0 if i in wc else (0.0 if i in wn else v) for i, v in zip(cc.id, yv)])
    res['auc_worst_case'] = round(float(fit_auc(worst)), 4)
    res['n_both_readers_against_primary'] = int(to_case.sum() + to_non.sum())

    tag = 'r2r3' + ('_adj' if adj_done else '_preadj')
    json.dump(res, open(os.path.join(REPO_OT, 'layer2_text_mode_agreement_%s.json' % tag), 'w'),
              ensure_ascii=False, indent=2)
    both32 = (a2[0] == 1) & (a3[0] == 1)
    mtx = pd.crosstab(pd.Series(a2[2][both32], name='R2'), pd.Series(a3[2][both32], name='R3'))
    mtx.to_csv(os.path.join(REPO_OT, 'layer2_text_mode_5level_matrix_%s.csv' % tag))
    if not quiet:
        print(json.dumps(res, ensure_ascii=False, indent=2))
        print("\n=== 五级分歧矩阵：R2（行）× R3（列）===")
        print(mtx.to_string())
    return res


# ---------------------------------------------------------------- selftest
def selftest():
    """端到端自检：模拟 R2/R3（不同错误率）+ R4 裁决，跑通 package 产物→adjudicate→analyse。"""
    tmp = r"E:/workbuddy/2026-09-29-18-58-07/tmpwork"
    os.makedirs(tmp, exist_ok=True)
    key = master_key()
    rng = np.random.default_rng(11)

    def simulate(rate_flip, rate_shift, rate_ng):
        g = key.gradeable.astype(int).values.copy()
        y = key.any_dr.fillna(0).astype(int).values.copy()
        s = key.severity.fillna(-1).values.copy()
        odd = (rng.random(len(key)) < rate_ng) & (g == 0)
        g[odd] = 1; s[odd] = 0.0; y[odd] = 0
        f = (rng.random(len(key)) < rate_flip) & (g == 1)
        y[f] = 1 - y[f]
        sh = (rng.random(len(key)) < rate_shift) & (g == 1)
        s[sh] = np.clip(s[sh] + rng.choice([-1, 1], size=int(sh.sum())), 0, 4)
        return dict(blind_id=key.blind_id,
                    grade=[np.nan if v < 0 else v for v in s], any=y, grad=g)

    def write(d, path):
        pd.DataFrame({
            'Blind ID': d['blind_id'],
            'Grade LEFT (0-4)': d['grade'], 'Grade RIGHT (0-4)': d['grade'],
            'Worst-eye grade (0-4)': ['' if g == 0 else v for g, v in zip(d['grad'], d['grade'])],
            'Any DR (1=yes, 0=no)': ['' if g == 0 else v for g, v in zip(d['grad'], d['any'])],
            'Gradable? (1=yes, 0=no)': d['grad'],
        }).to_excel(path, sheet_name=SHEET2, index=False)

    f2 = os.path.join(tmp, '_selftest_R2.xlsx')
    f3 = os.path.join(tmp, '_selftest_R3.xlsx')
    write(simulate(0.05, 0.02, 0.010), f2)
    write(simulate(0.07, 0.03, 0.005), f3)
    print("自检①：已模拟 R2、R3 两份回收件（错误率不同）")
    adj_path = build_adjudication(f2, f3)
    print("自检②：分歧表已生成")
    adj = pd.read_excel(adj_path, sheet_name=SHEET_ADJ)
    truth = key.set_index('blind_id')
    adj['ADJUDICATED gradable (1/0)'] = [int(truth.loc[b, 'gradeable']) for b in adj['Blind ID']]
    adj['ADJUDICATED grade (0-4)'] = ['' if truth.loc[b, 'gradeable'] == 0
                                      else float(truth.loc[b, 'severity']) for b in adj['Blind ID']]
    adj['ADJUDICATED any-DR (1/0)'] = ['' if truth.loc[b, 'gradeable'] == 0
                                       else float(truth.loc[b, 'any_dr']) for b in adj['Blind ID']]
    adj.to_excel(adj_path, sheet_name=SHEET_ADJ, index=False)
    print("自检③：已模拟 R4 裁决")
    res = analyse(f2, f3, adj_path)
    ok = (res.get('auc_with_consensus') is not None
          and not np.isnan(res['auc_with_consensus'])
          and res['n_discordant_requiring_adjudication'] > 0)
    print("\n自检结果：", "PASS" if ok else "FAIL")
    print("关键核验：主分析 AUC = %.4f（稿件报告 0.598）；共识后 AUC = %.4f；"
          "R2–R3 五级加权 κ = %.4f；R2–R3 any-DR κ = %.4f"
          % (res['auc_primary'], res['auc_with_consensus'],
             res['R2_vs_R3_inter_reader']['five_level_wkappa'],
             res['R2_vs_R3_inter_reader']['anydr_kappa']))
    return res


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('cmd', choices=['package', 'adjudicate', 'analyse', 'selftest'])
    ap.add_argument('--r2'); ap.add_argument('--r3'); ap.add_argument('--adj', default=None)
    a = ap.parse_args()
    if a.cmd == 'package':
        build_package()
    elif a.cmd == 'adjudicate':
        build_adjudication(a.r2, a.r3)
    elif a.cmd == 'selftest':
        selftest()
    else:
        if not (a.r2 and a.r3):
            raise SystemExit("analyse 需要 --r2 与 --r3")
        analyse(a.r2, a.r3, a.adj)
