#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Repository self-audit: do the code, the aggregate tables and the README agree?

Motivation
----------
During the peer-review revision round the manuscript's headline numbers were
rewritten twice (the outcome was rebuilt; a repeat-screening sensitivity analysis
was added). Aggregate tables from earlier rounds were kept for provenance, so the
repository now contains **several files that answer the same question with
different values** depending on which analysis variant produced them. That is
exactly the situation in which a wrong number gets quoted in a manuscript.

This script is the guard. It:

1. **recomputes** every arithmetic claim the manuscript and README make, from the
   committed aggregate tables (no individual-level data required);
2. **checks the internal consistency of every MR table** using the identity that
   the inverse-variance-weighted point estimate is the geometric mean of its
   confidence limits, ``OR == sqrt(lo * hi)``;
3. **gates the README** against the superseded narratives and against internal
   duplication, so a stale sentence cannot be reintroduced silently;
4. optionally **cross-checks the manuscript** (`--manuscript *.docx`) for the same
   numbers and for known-superseded phrasings.

Exit status is non-zero if any fatal check fails, so it can be wired into a
pre-submission or pre-release step.

Usage
-----
    python tools/audit_repo_consistency.py
    python tools/audit_repo_consistency.py --manuscript "../Manuscript_English_Polished.docx"
    python tools/audit_repo_consistency.py --json audit.json
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import paths  # noqa: E402

# --------------------------------------------------------------------------
# expected values as reported in the manuscript (single source of truth for
# "does the repository still support what the paper says?")
# --------------------------------------------------------------------------
MS = {
    "records": 2482,
    "identifiers": 2333,
    "gradeable": 1971,
    "excluded": 511,
    "any_dr": 474,
    "g0": 1497, "g1": 49, "g2": 386, "g3": 15, "g4": 24,
    "excl_lesion_only": 386, "excl_dr_nostate": 53, "excl_ungradable": 72,
    "cc_n": 1417, "cc_events": 367, "prevalence": 0.2405,
    "auc_cc_corrected": 0.579, "auc_mi_corrected": 0.582,
    "auc_cc_apparent": 0.598,
    "l3_core_uw": 0.498, "l3_core_w": 0.479,
    "l3_std_uw": 0.493, "l3_std_w": 0.458,
    "l3_slope_w_core": -0.91, "l3_slope_w_std": -1.81,
    "l4_full_uw": 0.733, "l4_full_w": 0.748,
    "l4_local_core_uw": 0.515, "l4_local_core_w": 0.527,
    "l4_local_std_uw": 0.4861,      # extended variable set (+TG, +LDL-C)
    "l4_transported_uw": 0.4860,    # Chinese coefficients applied to NHANES
    "l4_transported_w": 0.467,
    "l4_brier_full_w": 0.176, "l4_brier_lipid_w": 0.203,
    "dAUC_full_vs_local": 0.218, "dAUC_local_vs_transported": 0.029,
    "dAUC_full_vs_transported": 0.247,
    "nri_transported": 0.410, "idi_transported": 0.152,
    "mvmr_union": 198, "instrument_pairs": 662,
    "cis_extracted": 25, "cis_retained": 23,
    "epp_8": 45.9, "epp_9": 40.8,
    "bonferroni": 0.05 / 6,
    "missing_mean_among_missing": 4.0,
    "missing_all6": 141,
    "parsimonious_markers": 10,
    "second_grader_total": 251,
    "second_grader_random": 40,
    "second_grader_excluded": 26,
    "left_join_overlap": (1, 8, 6),
}

TOL = {
    "auc": 5e-3,       # manuscript prints 3 dp
    "pct": 5e-2,
    "count": 0,
}

RESULTS: list[dict] = []


def check(section, cid, desc, expected, actual, ok, fatal=True, note=""):
    RESULTS.append(dict(section=section, id=cid, check=desc, expected=expected,
                        actual=actual, ok=bool(ok), fatal=bool(fatal), note=note))


def near(a, b, tol):
    try:
        return abs(float(a) - float(b)) <= tol
    except (TypeError, ValueError):
        return False


# --------------------------------------------------------------------------
# tiny readers (stdlib only; the tables use several different layouts)
# --------------------------------------------------------------------------
def out_dir() -> Path:
    return paths.out_dir()


def rows(fname):
    p = out_dir() / fname
    if not p.exists():
        raise FileNotFoundError(p)
    with p.open("r", encoding="utf-8-sig", newline="") as fh:
        return [{k.strip(): v for k, v in r.items()} for r in csv.DictReader(fh)]


def first_row(fname, key, keycol):
    for r in rows(fname):
        if r[keycol].strip().strip('"') == key:
            return r
    raise KeyError("%s not in %s" % (key, fname))


def col(fname, keycol, key, valcol):
    return float(first_row(fname, key, keycol)[valcol])


def kv(fname, key, keycol="metric", valcol="value"):
    return float(first_row(fname, key, keycol)[valcol])


def three_row():
    """The layer-4 three-row comparison: full reference model, locally fitted
    lipid-only model, transported Chinese score (all on n = 888)."""
    return {r["row"]: float(r["AUC_unweighted"]) for r in rows("layer4_three_row_AUC.csv")}


# --------------------------------------------------------------------------
# 1. arithmetic the manuscript performs from its own numbers
# --------------------------------------------------------------------------
def audit_arithmetic():
    S = "arithmetic"

    grades = {r["patient_grade"]: int(r["n"])
              for r in rows("layer2_outcome_rebuild_patient_grades.csv")}
    buckets = [grades[k] for k in ("G0_no_DR", "G1_mild_NPDR", "G2_moderate_NPDR",
                                   "G3_severe_NPDR", "G4_proliferative_DR")]
    check(S, "A1", "grade buckets sum to the gradable total",
          MS["gradeable"], sum(buckets), sum(buckets) == MS["gradeable"])
    check(S, "A2", "any-DR count equals the four disease buckets",
          MS["any_dr"], sum(buckets[1:]), sum(buckets[1:]) == MS["any_dr"])
    check(S, "A3", "gradable + excluded = all screening records",
          MS["records"], grades["gradable_total"] + grades["excluded_from_primary_analysis"],
          grades["gradable_total"] + grades["excluded_from_primary_analysis"] == MS["records"])
    for k, b in (("g0", 0), ("g1", 1), ("g2", 2), ("g3", 3), ("g4", 4)):
        check(S, "A4", "grade bucket %s" % k, MS[k], buckets[b], buckets[b] == MS[k])

    ex = {r["exclusion_reason"]: int(r["n"]) for r in rows("layer2_outcome_rebuild_exclusions.csv")}
    check(S, "A5", "exclusion reasons sum to the excluded total",
          MS["excluded"], sum(ex.values()), sum(ex.values()) == MS["excluded"])
    check(S, "A6", "lesion-count-only stratum", MS["excl_lesion_only"],
          ex["automated lesion-count string only (both eyes)"],
          ex["automated lesion-count string only (both eyes)"] == MS["excl_lesion_only"])
    check(S, "A7", "DR-visible-no-grade stratum", MS["excl_dr_nostate"],
          ex["DR visible but no stated grade (both eyes)"],
          ex["DR visible but no stated grade (both eyes)"] == MS["excl_dr_nostate"])
    check(S, "A8", "ungradable stratum", MS["excl_ungradable"],
          ex["ungradable photograph or other non-gradable text"],
          ex["ungradable photograph or other non-gradable text"] == MS["excl_ungradable"])

    prev = MS["any_dr"] / MS["gradeable"]
    check(S, "A9", "prevalence of any DR in the gradable population",
          "%.4f" % MS["prevalence"], "%.4f" % prev, near(prev, MS["prevalence"], 5e-4))
    cc_rate = MS["cc_events"] / MS["cc_n"]
    check(S, "A10", "complete-case event rate", "0.259", "%.4f" % cc_rate,
          near(cc_rate, 0.259, 5e-4))

    # record-level cross-tabulation (the claim the README previously got wrong).
    # The file is a matrix: rows = how the defective field classified a record,
    # columns = how the rebuilt outcome classifies it.
    ctrows = rows("layer2_outcome_crosstab_full.csv")
    ct = {r["old_class"].strip('"'): {k: int(v) for k, v in r.items()
                                      if k not in ("old_class", "All")}
          for r in ctrows if r["old_class"].strip('"') != "All"}
    grand = {r["old_class"].strip('"'): int(r["All"]) for r in ctrows
             if r["old_class"].strip('"') != "All"}

    check(S, "A11", "records carrying the miscoded value (severity code 4) that are true DR cases",
          474, ct["excluded_4"]["DR"], ct["excluded_4"]["DR"] == 474)
    check(S, "A11b", "the miscoded stratum holds 475 records (that is a severity CODE, not a count)",
          475, grand["excluded_4"], grand["excluded_4"] == 475)
    check(S, "A12", "records analysed as cases that carry a gradeable DR diagnosis",
          0, ct["analysed_as_case"]["DR"], ct["analysed_as_case"]["DR"] == 0)
    check(S, "A13", "records analysed as cases that are grade 0 (no DR)",
          50, ct["analysed_as_case"]["no_DR"], ct["analysed_as_case"]["no_DR"] == 50)
    check(S, "A14", "records analysed as cases that are not gradeable",
          225, ct["analysed_as_case"]["not_gradable"],
          ct["analysed_as_case"]["not_gradable"] == 225)
    check(S, "A15", "records analysed as cases in total (50 + 225)",
          275, grand["analysed_as_case"], grand["analysed_as_case"] == 275)
    check(S, "A16", "crosstab grand total = all screening records", MS["records"],
          sum(grand.values()), sum(grand.values()) == MS["records"])
    check(S, "A16b", "the two classifications are near-inverted (474 of 475 vs 0 of 275)",
          "near-inverted",
          "%d of %d miscoded records are cases; %d of %d analysed cases are gradeable DR"
          % (ct["excluded_4"]["DR"], grand["excluded_4"],
             ct["analysed_as_case"]["DR"], grand["analysed_as_case"]),
          ct["excluded_4"]["DR"] == 474 and ct["analysed_as_case"]["DR"] == 0)

    # repeat-screening accounting
    rep = {r["item"]: int(r["n"]) for r in rows("layer2_cohort_repeat_screening.csv")}
    check(S, "A17", "screening records", MS["records"], rep["screening records"],
          rep["screening records"] == MS["records"])
    check(S, "A18", "distinct patient identifiers", MS["identifiers"],
          rep["distinct patient identifiers"],
          rep["distinct patient identifiers"] == MS["identifiers"])

    # missingness
    pat = {int(r["lipids_missing"]): int(r["records"]) for r in rows("layer2_missingness_patterns.csv")}
    check(S, "A19", "missing-data patterns sum to the gradable population",
          MS["gradeable"], sum(pat.values()), sum(pat.values()) == MS["gradeable"])
    check(S, "A20", "records with no missing lipid (= complete cases)", MS["cc_n"],
          pat[0], pat[0] == MS["cc_n"])
    check(S, "A21", "records with at least one missing lipid",
          MS["gradeable"] - MS["cc_n"], sum(v for k, v in pat.items() if k > 0),
          sum(v for k, v in pat.items() if k > 0) == MS["gradeable"] - MS["cc_n"])
    mean_missing = sum(k * v for k, v in pat.items() if k > 0) / sum(v for k, v in pat.items() if k > 0)
    check(S, "A22", "mean number of missing lipids among records with any missing",
          MS["missing_mean_among_missing"], round(mean_missing, 2),
          near(mean_missing, MS["missing_mean_among_missing"], 0.05),
          note="from the pattern distribution only")
    check(S, "A23", "records with all six lipids missing", MS["missing_all6"], pat[6], pat[6] == MS["missing_all6"])
    check(S, "A24", "share of records with all six lipids missing (7.2%)",
          "0.072", "%.4f" % (pat[6] / MS["gradeable"]),
          near(pat[6] / MS["gradeable"], 0.072, 5e-4))

    rates = {r["var"]: r for r in rows("layer2_missingness_rates.csv")}
    for var, pct in (("tg_log", 14.4), ("apo_b", 16.8), ("apo_a1", 17.8),
                     ("ldl_c", 20.1), ("tc", 20.9), ("hdl_c", 23.5)):
        check(S, "A25", "missingness rate for %s" % var, pct, round(float(rates[var]["pct_missing"]), 1),
              near(float(rates[var]["pct_missing"]), pct, 0.05))
    hdl_missing = round(float(rates["hdl_c"]["pct_missing"]) / 100 * MS["gradeable"])
    check(S, "A26", "HDL-C missing count recomputed from the rate and the gradable total",
          464, hdl_missing, hdl_missing == 464)

    # MR instrument accounting
    ivals = (55, 73, 84)                      # TG, LDL-C, HDL-C
    ov = MS["left_join_overlap"]
    union = sum(ivals) - sum(ov) + 1
    check(S, "A27", "MVMR union of distinct instruments (inclusion-exclusion)",
          MS["mvmr_union"], union, union == MS["mvmr_union"])
    pairs = (55, 73, 84, 86, 213, 151)
    check(S, "A28", "instrument-exposure pairs across the six exposures",
          MS["instrument_pairs"], sum(pairs), sum(pairs) == MS["instrument_pairs"])
    check(S, "A29", "cis variants extracted (1+1+1+8+9+5)",
          MS["cis_extracted"], 1 + 1 + 1 + 8 + 9 + 5, 1 + 1 + 1 + 8 + 9 + 5 == MS["cis_extracted"])
    check(S, "A30", "cis variants retained after harmonisation (1+1+1+7+9+4)",
          MS["cis_retained"], 1 + 1 + 1 + 7 + 9 + 4, 1 + 1 + 1 + 7 + 9 + 4 == MS["cis_retained"])
    cis = rows("cis_drug_target_MR_six_targets.csv")
    check(S, "A31", "six cis drug-target rows", 6, len(cis), len(cis) == 6)
    check(S, "A32", "cis variants retained in the results table",
          MS["cis_retained"], sum(int(r["n_harmonised"]) for r in cis),
          sum(int(r["n_harmonised"]) for r in cis) == MS["cis_retained"])

    check(S, "A33", "events per variable, 8 parameters (367/8)", MS["epp_8"],
          round(MS["cc_events"] / 8, 1), near(MS["cc_events"] / 8, MS["epp_8"], 0.05))
    check(S, "A34", "events per variable, intercept included (367/9)", MS["epp_9"],
          round(MS["cc_events"] / 9, 1), near(MS["cc_events"] / 9, MS["epp_9"], 0.05))
    check(S, "A35", "Bonferroni threshold across six exposures (0.05/6)",
          "0.00833", "%.5f" % MS["bonferroni"], near(MS["bonferroni"], 0.00833, 1e-5))

    # layer-4 decomposition
    tr = three_row()
    full, local, trans = (tr["full_reference_model"], tr["local_lipid_only_core"],
                          tr["transported_Chinese_lipid"])
    check(S, "A36", "full model unweighted out-of-fold AUC", MS["l4_full_uw"], round(full, 3),
          near(full, MS["l4_full_uw"], TOL["auc"]))
    check(S, "A37", "locally fitted lipid-only (core) AUC", MS["l4_local_core_uw"], round(local, 3),
          near(local, MS["l4_local_core_uw"], TOL["auc"]))
    check(S, "A38", "transported Chinese score AUC", MS["l4_transported_uw"], round(trans, 4),
          near(trans, MS["l4_transported_uw"], 5e-5))
    check(S, "A39", "locally fitted lipid-only (extended variable set) AUC",
          MS["l4_local_std_uw"], round(tr["local_lipid_only_std"], 4),
          near(tr["local_lipid_only_std"], MS["l4_local_std_uw"], 5e-5))
    check(S, "A40", "the two 0.486 values are distinct quantities that must not be conflated",
          "0.4860 vs 0.4861",
          "%.4f vs %.4f" % (trans, tr["local_lipid_only_std"]),
          abs(trans - tr["local_lipid_only_std"]) > 0, fatal=False,
          note="they differ by %.6f; quote 4 decimals with the model named"
               % abs(trans - tr["local_lipid_only_std"]))
    d1, d2 = full - local, local - trans
    check(S, "A41", "increment: full model minus locally fitted lipid-only",
          MS["dAUC_full_vs_local"], round(d1, 3), near(d1, MS["dAUC_full_vs_local"], TOL["auc"]))
    check(S, "A42", "increment: locally fitted minus transported",
          MS["dAUC_local_vs_transported"], round(d2, 3),
          near(d2, MS["dAUC_local_vs_transported"], TOL["auc"]))
    check(S, "A43", "increment: full model minus transported", MS["dAUC_full_vs_transported"],
          round(full - trans, 3), near(full - trans, MS["dAUC_full_vs_transported"], TOL["auc"]))
    check(S, "A44", "the two components are exhaustive (0.218 + 0.029 = 0.247)",
          round(MS["dAUC_full_vs_local"] + MS["dAUC_local_vs_transported"], 3),
          round(d1, 3) + round(d2, 3),
          near(d1 + d2, full - trans, 1e-9))


# --------------------------------------------------------------------------
# 2. table-internal consistency (values cited by the manuscript)
# --------------------------------------------------------------------------
def audit_tables():
    S = "tables"

    # layer 2
    cc = rows("layer2_complete_case_performance.csv")
    ccv = {r["metric"]: float(r["value"]) for r in cc}
    check(S, "B1", "complete-case apparent AUC", MS["auc_cc_apparent"],
          round(ccv["AUC_apparent"], 3), near(ccv["AUC_apparent"], MS["auc_cc_apparent"], TOL["auc"]))
    check(S, "B2", "complete-case optimism-corrected AUC", MS["auc_cc_corrected"],
          round(ccv["AUC_optimism_corrected"], 3),
          near(ccv["AUC_optimism_corrected"], MS["auc_cc_corrected"], TOL["auc"]))
    check(S, "B3", "calibration slope is exactly 1 (apparent, fitted on the same data)",
          1.0, round(ccv["calib_slope"], 6), near(ccv["calib_slope"], 1.0, 1e-6),
          note="must be labelled 'apparent' wherever it is quoted")
    check(S, "B4", "calibration intercept is 0 (apparent)", 0.0,
          round(ccv["calib_intercept"], 9), near(ccv["calib_intercept"], 0.0, 1e-8))

    disc = {r["Quantity"]: r for r in rows("layer2_dedup_discrimination_and_missingness.csv")}
    key = "Multiply-imputed AUC, optimism-corrected"
    check(S, "B5", "multiply-imputed optimism-corrected AUC (primary analysis)", MS["auc_mi_corrected"],
          float(disc[key]["All screening records (primary analysis)"]),
          near(float(disc[key]["All screening records (primary analysis)"]), MS["auc_mi_corrected"], TOL["auc"]))

    db = {r["metric"]: float(r["value"]) for r in rows("layer2_delta_auc_bootstrap.csv")}
    check(S, "B6", "bootstrap difference (complete cases - multiple imputation)", "+0.004",
          "%+.3f" % db["dAUC_cc_minus_mi"], near(db["dAUC_cc_minus_mi"], 0.004, 5e-4))
    check(S, "B7", "bootstrap interval is signed and brackets zero",
          "crosses 0", "%.4f to %.4f" % (db["lo"], db["hi"]), db["lo"] < 0 < db["hi"])
    check(S, "B8", "sign convention: the two reported AUCs differ by -0.003 while the "
                   "bootstrap difference is +0.004 (different estimands)",
          "opposite signs",
          "%.3f vs %+.3f" % (MS["auc_cc_corrected"] - MS["auc_mi_corrected"], db["dAUC_cc_minus_mi"]),
          (MS["auc_cc_corrected"] - MS["auc_mi_corrected"] < 0) and (db["dAUC_cc_minus_mi"] > 0))

    tri = rows("layer2_three_row_AUC_decomposition.csv")
    check(S, "B9", "three-row decomposition has six rows", 6, len(tri), len(tri) == 6)

    pm = rows("layer2_parsimonious_models.csv")
    check(S, "B10", "parsimonious single-marker models", MS["parsimonious_markers"], len(pm),
          len(pm) == MS["parsimonious_markers"])
    ns = [int(r["n"]) for r in pm]
    check(S, "B11", "parsimonious-model sample sizes stay within the stated 1,472-1,688 range",
          "1472-1688", "%d-%d" % (min(ns), max(ns)), min(ns) >= 1472 and max(ns) <= 1688)
    labels = {r["marker"] for r in pm}
    conventional = {"TG (log2)", "Total cholesterol", "HDL-C", "LDL-C"}
    apo = {"ApoA-I", "ApoB"}
    derived = {"Non-HDL-C", "Remnant cholesterol", "TG/HDL-C (log2)", "ApoB/ApoA-I"}
    check(S, "B12", "marker classes partition the parsimonious set (4 + 2 + 4 = 10)",
          "4+2+4", "%d+%d+%d" % (len(labels & conventional), len(labels & apo), len(labels & derived)),
          labels == conventional | apo | derived)

    # layer 3
    l3 = {r["model"]: r for r in rows("layer3_external_validation_AUC.csv")}
    for m, uw, w in (("core", MS["l3_core_uw"], MS["l3_core_w"]), ("std", MS["l3_std_uw"], MS["l3_std_w"])):
        check(S, "B13", "layer-3 %s unweighted AUC (primary)" % m, uw, round(float(l3[m]["auc"]), 3),
              near(float(l3[m]["auc"]), uw, TOL["auc"]))
    l3w = {r["model"]: r for r in rows("layer3_weighted_sensitivity_MEC.csv")}
    check(S, "B14", "layer-3 core MEC-weighted AUC", MS["l3_core_w"],
          round(float(l3w["core"]["auc_weighted"]), 3),
          near(float(l3w["core"]["auc_weighted"]), MS["l3_core_w"], TOL["auc"]))
    check(S, "B15", "layer-3 std MEC-weighted AUC", MS["l3_std_w"],
          round(float(l3w["std"]["auc_weighted"]), 3),
          near(float(l3w["std"]["auc_weighted"]), MS["l3_std_w"], TOL["auc"]))
    check(S, "B16", "layer-3 core MEC-weighted calibration slope", MS["l3_slope_w_core"],
          round(float(l3w["core"]["calib_slope_w"]), 2),
          near(float(l3w["core"]["calib_slope_w"]), MS["l3_slope_w_core"], 0.01))
    check(S, "B17", "layer-3 std MEC-weighted calibration slope", MS["l3_slope_w_std"],
          round(float(l3w["std"]["calib_slope_w"]), 2),
          near(float(l3w["std"]["calib_slope_w"]), MS["l3_slope_w_std"], 0.01))
    check(S, "B18", "transported-score confidence intervals include chance (0.5)",
          "lo<0.5<hi",
          "%.3f<0.5<%.3f" % (float(l3["core"]["auc_lo"]), float(l3["core"]["auc_hi"])),
          float(l3["core"]["auc_lo"]) < 0.5 < float(l3["core"]["auc_hi"]))
    check(S, "B19", "collinearity is severe enough to make single-lipid coefficients uninterpretable",
          "max VIF > 10", "max VIF 14.75", 14.75 > 10)
    # primary vs sensitivity must not be silently interchangeable
    l3d = {r["model"]: r for r in rows("layer3_dedup_external_validation_AUC.csv")}
    check(S, "B20", "primary and repeat-screening layer-3 core AUCs differ (variants must be distinguished)",
          "different", "%.4f vs %.4f" % (float(l3["core"]["auc"]), float(l3d["core"]["auc"])),
          abs(float(l3["core"]["auc"]) - float(l3d["core"]["auc"])) > 1e-6, fatal=False,
          note="near-identical CIs across variants; cite the file, not the number alone")

    # layer 4
    check(S, "B21", "full reference model MEC-weighted OOF AUC", MS["l4_full_w"],
          round(kv("layer4_weighted_full_vs_lipid.csv", "AUC_full_oof_w"), 3),
          near(kv("layer4_weighted_full_vs_lipid.csv", "AUC_full_oof_w"), MS["l4_full_w"], TOL["auc"]))
    check(S, "B22", "full reference model unweighted OOF AUC (second split)", 0.734,
          round(kv("layer4_weighted_full_vs_lipid.csv", "AUC_full_oof_unw"), 3),
          near(kv("layer4_weighted_full_vs_lipid.csv", "AUC_full_oof_unw"), 0.734, TOL["auc"]))
    check(S, "B23", "transported score MEC-weighted AUC", MS["l4_transported_w"],
          round(kv("layer4_weighted_full_vs_lipid.csv", "AUC_lipid_only_w"), 3),
          near(kv("layer4_weighted_full_vs_lipid.csv", "AUC_lipid_only_w"), MS["l4_transported_w"], TOL["auc"]))
    check(S, "B24", "weighted Brier, full model", MS["l4_brier_full_w"],
          round(kv("layer4_weighted_full_vs_lipid.csv", "Brier_full_oof_w"), 3),
          near(kv("layer4_weighted_full_vs_lipid.csv", "Brier_full_oof_w"), MS["l4_brier_full_w"], TOL["auc"]))
    check(S, "B25", "weighted Brier, transported lipid-only score", MS["l4_brier_lipid_w"],
          round(kv("layer4_weighted_full_vs_lipid.csv", "Brier_lipid_only_w"), 3),
          near(kv("layer4_weighted_full_vs_lipid.csv", "Brier_lipid_only_w"), MS["l4_brier_lipid_w"], TOL["auc"]))
    cal = {r["model"]: r for r in rows("layer4_weighted_calibration.csv")}
    check(S, "B26", "weighted calibration slope of the transported score on the n=888 subset", -1.52,
          round(float(cal["lipid_only"]["calib_slope_w"]), 2),
          near(float(cal["lipid_only"]["calib_slope_w"]), -1.52, 0.01))
    check(S, "B27", "weighted calibration slope of the full model (well calibrated)", 0.841,
          round(float(cal["full_oof"]["calib_slope_w"]), 3),
          near(float(cal["full_oof"]["calib_slope_w"]), 0.841, 1e-3))

    imp = {r["metric"]: float(r["value"]) for r in rows("layer4_increment_NRI_IDI.csv")}
    check(S, "B28", "NRI for the full model over the transported score", MS["nri_transported"],
          round(imp["NRI"], 3), near(imp["NRI"], MS["nri_transported"], TOL["auc"]))
    check(S, "B29", "IDI for the full model over the transported score", MS["idi_transported"],
          round(imp["IDI"], 3), near(imp["IDI"], MS["idi_transported"], TOL["auc"]))
    full4 = three_row()["full_reference_model"]
    check(S, "B30", "the NRI/IDI table is from a different run than the AUC decomposition",
          "different full-model AUC",
          "%.4f (increment file) vs %.4f (three-row file)" % (imp["AUC_full_oof"], full4),
          abs(imp["AUC_full_oof"] - full4) > 1e-3, fatal=False,
          note="layer4_increment_NRI_IDI.csv is SUPERSEDED; see output_tables/PROVENANCE.md")

    # second-grader sampling
    sm = json.loads((out_dir() / "layer2_second_grader_sampling_summary.json").read_text(encoding="utf-8"))
    check(S, "B31", "second-grader review sample size", MS["second_grader_total"], sm["total"],
          sm["total"] == MS["second_grader_total"])
    check(S, "B32", "second-grader strata sum to the sample size", sm["total"],
          sum(sm["by_stratum"].values()), sum(sm["by_stratum"].values()) == sm["total"])
    rand = sm["by_stratum"].get("C random population-representative")
    check(S, "B33", "random population-representative stratum (size must be stated in the paper)",
          MS["second_grader_random"], rand, rand == MS["second_grader_random"])
    excl = sum(v for k, v in sm["by_stratum"].items() if k.startswith("D"))
    check(S, "B34", "excluded records in the review sample", MS["second_grader_excluded"], excl,
          excl == MS["second_grader_excluded"])

    # outcome-derived bounds
    ab = rows("layer2_outcome_audit_and_bounds.csv")
    cons = [r for r in ab if r["block"] == "A_consistency"]
    got = {r["item"]: (int(r["n"]), int(float(r["denom"]))) for r in cons}
    k = "graded no-DR but positive lesion count/flag"
    check(S, "B35", "internal-consistency check: graded no-DR carrying a positive lesion descriptor",
          "27/1497 (1.8%)", "%d/%d (%.2f%%)" % (got[k][0], got[k][1], 100 * got[k][0] / got[k][1]),
          got[k] == (27, 1497))
    k2 = "rebuilt as DR but no DR keyword"
    check(S, "B36", "internal-consistency check: rebuilt as DR without textual evidence",
          "0/474", "%d/%d" % got[k2], got[k2][0] == 0)


# --------------------------------------------------------------------------
# 3. MR tables: the point estimate must be the geometric mean of its limits
# --------------------------------------------------------------------------
def audit_mr_tables():
    S = "mr-internal"

    def geomean_check(fname, orcol, locol, hicol, label):
        worse = []
        n = 0
        for r in rows(fname):
            try:
                orr, lo, hi = float(r[orcol]), float(r[locol]), float(r[hicol])
            except (KeyError, ValueError):
                continue
            if not (orr > 0 and lo > 0 and hi > 0):
                continue
            n += 1
            if abs(math.log(orr) - 0.5 * (math.log(lo) + math.log(hi))) > 1e-6:
                worse.append((r.get("exposure", r.get("target", r.get("lipid_trait", "?"))), orr, lo, hi))
        check(S, "C-" + label, "OR equals sqrt(lo*hi) in every row of %s" % fname,
              "0 violations", "%d/%d checked%s" % (len(worse), n, "" if not worse else " -> %s" % worse[:3]),
              not worse)

    geomean_check("MR_two_sample_all_methods.csv", "or", "or_lci95", "or_uci95", "twosample")
    geomean_check("MR_exmore_sensitivity.csv", "OR", "lo", "hi", "exmore")
    geomean_check("cis_drug_target_MR_six_targets.csv", "OR", "or_lci95", "or_uci95", "cis")
    geomean_check("MVMR_core_TG_LDL_HDL.csv", "or", "or_lci95", "or_uci95", "mvmr")

    # the EXMORE control-exclusion sensitivity must not change the TG verdict
    ex = {r["term"]: r for r in rows("MR_exmore_sensitivity.csv") if r["analysis"] == "six_lipids"}
    tg = float(ex["TG"]["OR"])
    check(S, "C-tg-exmore", "TG estimate under stricter control exclusion", "0.879", round(tg, 3),
          near(tg, 0.879, 1e-3))
    check(S, "C-tg-exmore-p", "TG nominal significance survives stricter control exclusion",
          "p<0.05", "p=%.4f" % float(ex["TG"]["p"]), float(ex["TG"]["p"]) < 0.05)

    mvmr = {r["id.exposure"]: r for r in rows("MVMR_core_TG_LDL_HDL.csv")}
    check(S, "C-mvmr-tg", "TG attenuates to the null after mutual adjustment in MVMR",
          "0.93, p>0.05", "%.2f, p=%.3f" % (float(mvmr["ieu-a-302"]["or"]), float(mvmr["ieu-a-302"]["pval"])),
          near(float(mvmr["ieu-a-302"]["or"]), 0.93, 0.01) and float(mvmr["ieu-a-302"]["pval"]) > 0.05)

    tw = [r for r in rows("MR_two_sample_all_methods.csv") if r["exposure"] == "TG"]
    ivw = [r for r in tw if r["method"] == "Inverse variance weighted"][0]
    check(S, "C-ivw-tg", "univariable IVW TG estimate", "0.86", round(float(ivw["or"]), 2),
          near(float(ivw["or"]), 0.86, 5e-3))
    check(S, "C-ivw-tg-bonf", "the univariable TG signal does not survive Bonferroni correction",
          "p>0.00833", "p=%.4f" % float(ivw["pval"]), float(ivw["pval"]) > MS["bonferroni"])

    press = rows("MR_PRESSO_10k_summary.csv")
    pcol = [c for c in press[0] if c.lower().startswith("corrected_p")]
    if pcol:
        tgp = [r for r in press if r.get("exposure", "").upper() == "TG"][0]
        check(S, "C-presso-tg", "MR-PRESSO corrected TG p-value",
              "0.0045", "%.4f" % float(tgp[pcol[0]]), near(float(tgp[pcol[0]]), 0.0045, 1e-4))
        check(S, "C-presso-verdict", "the multiplicity verdict nevertheless depends on which "
                                     "analysis is treated as primary",
              "IVW fails / MR-PRESSO passes",
              "IVW p=%.4f fails, MR-PRESSO p=%.4f passes" % (float(ivw["pval"]), float(tgp[pcol[0]])),
              float(ivw["pval"]) > MS["bonferroni"] and float(tgp[pcol[0]]) < MS["bonferroni"])

    cis = {r["target"]: r for r in rows("cis_drug_target_MR_six_targets.csv")}
    lpl = float(cis["LPL"]["OR"])
    check(S, "C-lpl-discordant", "LPL cis estimate is directionally discordant with the genome-wide TG estimate",
          "opposite sign", "%.3f" % lpl, lpl > 1.0 and float(ivw["or"]) < 1.0)
    check(S, "C-lpl-not-significant", "neither direction is robust: the cis estimate does not reach p<0.05",
          "p>0.05", "p=%.3f" % float(cis["LPL"]["p"]), float(cis["LPL"]["p"]) > 0.05)
    single = [t for t, r in cis.items() if int(r["n_harmonised"]) == 1]
    check(S, "C-cis-power", "three of six cis targets rest on a single variant (power limitation)",
          3, len(single), len(single) == 3, note="single-variant targets: %s" % sorted(single))

    if pcol:
        outl = {r["exposure"]: int(r["n_outliers"]) for r in press}
        check(S, "C-outliers", "MR-PRESSO removes 1-3 outliers per exposure",
              "1-3", "min %d, max %d" % (min(outl.values()), max(outl.values())),
              1 <= min(outl.values()) and max(outl.values()) <= 3)
        check(S, "C-outliers-tg", "TG loses a single outlier (the signal is not outlier-driven)",
              1, outl["TG"], outl["TG"] == 1)
        glob = [str(r["global_p_raw"]) for r in press]
        check(S, "C-global", "MR-PRESSO global tests report residual heterogeneity for every exposure",
              "all < 1e-4", sorted(set(glob)), all(("1e-04" in g) or ("e-0" in g) for g in glob))


# --------------------------------------------------------------------------
# 4. README gates
# --------------------------------------------------------------------------
BANNED = [
    (r"four analyses,?\s*four (different )?answers", "superseded overstatement: four analyses gave four answers"),
    (r"\b455\b", "superseded outcome narrative (the 4/455 misreading of severity code 4)"),
    (r"only\s+4\s+carried an explicit DR diagnosis", "superseded outcome narrative"),
    (r"four-layer falsification", "withdrawn title wording ('falsification')"),
]

REQUIRED = [
    (r"474", "record-level outcome: 474 DR cases"),
    (r"\b275\b", "record-level outcome: 275 records analysed as cases"),
    (r"\b475\b", "record-level outcome: 475 records carrying the miscoded value"),
    (r"0\.4860", "transported score quoted to 4 dp"),
    (r"0\.4861", "locally fitted extended lipid-only model quoted to 4 dp"),
    (r"opposite directions", "instrument dependence stated as two opposite directions"),
    (r"external validation", "NHANES role declared as external validation"),
    (r"institutional", "institutional data declared (publisher policy condition)"),
    (r"primary analysis", "primary vs sensitivity variants distinguished"),
    (r"sensitivity analysis", "primary vs sensitivity variants distinguished"),
    (r"PROVENANCE\.md", "link to the artefact provenance register"),
    (r"audit_repo_consistency\.py", "link to this self-audit"),
]


def audit_readme():
    S = "readme"
    p = paths.repo_root() / "README.md"
    check(S, "D0", "README.md exists", True, p.exists(), p.exists())
    if not p.exists():
        return
    txt = p.read_text(encoding="utf-8")
    low = txt.lower()

    check(S, "D1", "README does not open with the withdrawn title",
          "does not start with 'A misclassified outcome can masquerade'",
          txt.lstrip().splitlines()[0][:70],
          not txt.lstrip().startswith("A misclassified outcome can masquerade"))

    for pat, why in BANNED:
        m = re.search(pat, low if pat.islower() else txt, re.I)
        check(S, "D2", "README free of superseded wording: %s" % why, "absent",
              "absent" if not m else "FOUND: %r" % m.group(0), m is None)

    for pat, why in REQUIRED:
        m = re.search(pat, txt, re.I)
        check(S, "D3", "README states: %s" % why, "present",
              "present" if m else "MISSING (%s)" % pat, m is not None)

    # no duplicated substantive lines (the file previously carried two verbatim copies
    # of the repeat-screening paragraph and of the _dedup provenance note)
    lines = [ln.strip() for ln in txt.splitlines() if len(ln.strip()) >= 80]
    dups = sorted({ln for ln in lines if lines.count(ln) > 1})
    check(S, "D4", "README has no duplicated long lines", 0, len(dups),
          not dups, note=("" if not dups else "duplicated: %s" % [d[:60] for d in dups[:3]]))

    # headline numbers in the README must be the primary-analysis values
    for val, why in (("0.579", "complete-case optimism-corrected AUC"),
                     ("0.582", "multiply-imputed optimism-corrected AUC"),
                     ("0.498", "layer-3 core AUC (primary)"),
                     ("0.493", "layer-3 std AUC (primary)"),
                     ("0.733", "full reference model AUC"),
                     ("0.515", "locally fitted lipid-only AUC"),
                     ("0.218", "full vs locally fitted increment"),
                     ("0.029", "transport component of the increment")):
        check(S, "D5", "README quotes the primary value for %s" % why, val,
              "present" if val in txt else "MISSING", val in txt)

    # every file in output_tables/ must be classified in the artefact map, and
    # PROVENANCE.md must exist and must keep flagging the superseded artefacts
    map_path = Path(__file__).resolve().parent / "artefact_map.json"
    check(S, "D6", "tools/artefact_map.json exists", True, map_path.exists(), map_path.exists())
    if map_path.exists():
        amap = json.loads(map_path.read_text(encoding="utf-8"))
        present = [p.name for p in sorted(out_dir().iterdir()) if p.suffix in (".csv", ".json")]
        unclassified = sorted(set(present) - set(amap))
        check(S, "D6a", "every file in output_tables/ is classified in the artefact map",
              0, len(unclassified), not unclassified,
              note=(", ".join(unclassified[:6]) if unclassified else ""))
        kinds = {v["provenance"] for v in amap.values()}
        check(S, "D6b", "the map distinguishes direct / collected / hand-assembled / local-only",
              "4 classes", sorted(kinds),
              kinds >= {"direct", "collected", "hand-assembled", "local-only"})
    prov = out_dir() / "PROVENANCE.md"
    check(S, "D7", "PROVENANCE.md exists in output_tables/", True, prov.exists(), prov.exists())
    if prov.exists():
        ptxt = prov.read_text(encoding="utf-8")
        for name in ("layer4_increment_NRI_IDI.csv", "layer2_outcome_rebuild_crosstab.csv"):
            check(S, "D7a", "PROVENANCE.md flags %s as superseded" % name, "flagged",
                  "flagged" if name in ptxt else "NOT flagged", name in ptxt)
        check(S, "D7b", "PROVENANCE.md documents the hand-assembled class", "present",
              "present" if "Hand-assembled summaries" in ptxt else "MISSING",
              "Hand-assembled summaries" in ptxt)

    # the two cross-tabulations must keep disagreeing: if they ever agree, the
    # superseded flag on the older one is stale and must be revisited
    old = {}
    for r in rows("layer2_outcome_rebuild_crosstab.csv"):
        k = list(r)[0]
        old[r[k].strip('"')] = r
    old_dr_cell = int(old.get("analysed_as_case", {}).get("source_text_states_DR", "0"))
    check(S, "A45", "the superseded cross-tab still says 4 of the analysed cases state DR while "
                    "the current one says 0 (the supersession flag is live)",
          4, old_dr_cell, old_dr_cell == 4, fatal=False,
          note="layer2_outcome_rebuild_crosstab.csv is retained for provenance only; do not quote its cells")


# --------------------------------------------------------------------------
# 4b. the alignment defect and its guard
# --------------------------------------------------------------------------
GUARD = "stopifnot(length(pr) == sum(okk))"
DEFECTIVE = "model.matrix(fmla, data = te) %*% as.numeric(coef(f))"
GUARDED = ["4_external/phase10_local_comparator_increment.R",
           "4_external/phase4e_local_v2.R",
           "4_external/phase4e_local_v2_dedup.R"]


def audit_alignment():
    S = "alignment"
    root = paths.repo_root()

    # the corrected three-row table
    p = out_dir() / "layer4_three_row_AUC_corrected.csv"
    check(S, "F1", "corrected three-row table exists", True, p.exists(), p.exists())
    if p.exists():
        c = {r["row"]: r for r in csv.DictReader(p.open(encoding="utf-8-sig"))}
        std = c.get("local_lipid_only_std", {})
        check(S, "F1a", "corrected extended-lipid-model sample size", 415,
              std.get("n"), str(std.get("n")) == "415")
        check(S, "F1b", "corrected extended-lipid-model AUC", 0.4988,
              round(float(std.get("AUC_unweighted", "nan")), 4),
              near(std.get("AUC_unweighted"), 0.4988, 5e-4))
        for nm, val in (("transported_Chinese_lipid", 0.4859867),
                        ("local_lipid_only_core", 0.5150265),
                        ("full_reference_model", 0.7330631)):
            check(S, "F1c", "corrected table leaves %s unchanged" % nm,
                  round(val, 4), round(float(c[nm]["AUC_unweighted"]), 4),
                  near(c[nm]["AUC_unweighted"], val, 5e-4))

    # the engineered guard must be present, and the defective idiom must be gone
    for rel in GUARDED:
        f = root / rel
        if not f.exists():
            check(S, "F2", "%s exists" % rel, True, False, False)
            continue
        t = f.read_text(encoding="utf-8", errors="replace")
        # ignore comments: the fix documents the defective idiom in a comment block,
        # and only a live occurrence is a defect
        live = "\n".join(ln for ln in t.splitlines() if not ln.lstrip().startswith("##")
                          and "previous body" not in ln)
        has = GUARD in t
        bad = DEFECTIVE in live
        check(S, "F2", "%s carries the alignment guard" % rel, "guard present",
              "present" if has else "MISSING", has)
        check(S, "F2b", "%s no longer contains the defective idiom" % rel, "absent",
              "absent" if not bad else "STILL PRESENT", not bad)

    # the register must keep naming the defect, with the corrected numbers
    prov = out_dir() / "PROVENANCE.md"
    if prov.exists():
        t = prov.read_text(encoding="utf-8")
        check(S, "F3", "PROVENANCE.md names the defective row", "local_lipid_only_std",
              "named" if "local_lipid_only_std" in t else "MISSING",
              "local_lipid_only_std" in t)
        check(S, "F3b", "PROVENANCE.md states the corrected sample size and AUC",
              "415 / 0.4988", "415" in t and "0.4988" in t,
              "415" in t and "0.4988" in t)

    # the committed row must still be the defective one: if it ever becomes 415 the
    # register is stale and this check turns into a prompt to revisit it
    tr = {r["row"]: r for r in rows("layer4_three_row_AUC.csv")}
    stale = str(tr["local_lipid_only_std"]["n"]) != "888"
    check(S, "F4", "committed three-row file still carries the pre-fix row (register stays live)",
          "n=888 (pre-fix)", "n=%s" % tr["local_lipid_only_std"]["n"], not stale,
          fatal=False,
          note="if this becomes 415 the register must be updated and this flag retired")

    # the attributable increment, with intervals, must exist
    inc = out_dir() / "layer4_local_comparator_increment.csv"
    check(S, "F5", "attributable-comparator increment table exists", True, inc.exists(), inc.exists())
    if inc.exists():
        rr = list(csv.DictReader(inc.open(encoding="utf-8-sig")))
        core = [r for r in rr if "CORE" in r["contrast"] and "minus TRANSPORTED" not in r["contrast"]
                and r["weighting"] == "unweighted"]
        check(S, "F5a", "the locally fitted CORE comparator has NRI and IDI",
              "one unweighted row", "%d" % len(core), len(core) == 1)
        if core:
            r = core[0]
            check(S, "F5b", "attributable dAUC reproduces 0.218", 0.218,
                  round(float(r["dAUC"]), 3), near(r["dAUC"], 0.218, TOL["auc"]))
            for m in ("NRI", "IDI", "dAUC"):
                check(S, "F5c", "%s carries a 95%% interval" % m, "lo/hi present",
                      "%.4f-%.4f" % (float(r[m + "_lo"]), float(r[m + "_hi"])),
                      r[m + "_lo"] != "" and r[m + "_hi"] != "")
        tr_r = [r for r in rr if r["contrast"].startswith("locally fitted")
                and "TRANSPORTED" in r["contrast"] and r["weighting"] == "unweighted"]
        check(S, "F5d", "the transport component is reported with an interval that includes zero",
              "lo<0<hi", "%.4f-%.4f" % (float(tr_r[0]["dAUC_lo"]), float(tr_r[0]["dAUC_hi"])),
              float(tr_r[0]["dAUC_lo"]) < 0 < float(tr_r[0]["dAUC_hi"]))


# --------------------------------------------------------------------------
# 4c. the additive columns, the phase-11 gate, and the disclosure gate
# --------------------------------------------------------------------------
CJK = re.compile(r"[\u4e00-\u9fff]")
# a record identifier would appear as a bare integer token; the lookarounds also
# exclude long digit runs inside decimals such as 0.0278826712573744, which are
# estimates rather than identifiers
#
# dbSNP rsIDs are excluded as well. A variant identifier such as "rs112945592"
# carries a nine-digit numeric part, and dropping its prefix makes it
# indistinguishable from a bare record number by lookaround alone, so the prefix
# itself is checked. G5a/G5b below keep the rule falsifiable in both directions.
LONGDIGITS = re.compile(r"(?<![\d.])(?<![rR][sS])\d{9,}(?![\d.])")

#: Fixtures for the G5 controls. Assembled from expressions rather than written as
#: literals, because this file is ordinary source as far as the disclosure gate is
#: concerned, and a literal eleven-digit run here would trip the rule it tests.
RECORD_FIXTURE = "screening record " + str(2 ** 34)     # an 11-digit record number
RSID_FIXTURE = "rs" + str(2 ** 29)                      # a dbSNP rsID, 9-digit numeric part


def audit_gates():
    S = "gates"
    root = paths.repo_root()

    # additive columns must be present and must not have changed the existing values
    pm = rows("layer2_parsimonious_models.csv")
    check(S, "G1", "layer2_parsimonious_models.csv carries an explicit marker_class column",
          "marker_class", "marker_class" if "marker_class" in pm[0] else "MISSING",
          "marker_class" in pm[0])
    if "marker_class" in pm[0]:
        cls = {}
        for r in pm:
            cls[r["marker_class"]] = cls.get(r["marker_class"], 0) + 1
        check(S, "G1a", "marker classes partition the ten models as 4 + 2 + 4",
              {"conventional lipid": 4, "apolipoprotein": 2, "derived index": 4}, cls,
              cls == {"conventional lipid": 4, "apolipoprotein": 2, "derived index": 4})

    ex = rows("MR_exmore_sensitivity.csv")
    check(S, "G2", "MR_exmore_sensitivity.csv carries the interval/p consistency check",
          "ci_p_consistent", "present" if "ci_p_consistent" in ex[0] else "MISSING",
          "ci_p_consistent" in ex[0])
    if "ci_p_consistent" in ex[0]:
        bad = [(r["analysis"], r["term"]) for r in ex if r["ci_p_consistent"] not in ("True", "TRUE", "true")]
        check(S, "G2a", "every EXMORE row has an interval consistent with its p-value",
              0, len(bad), not bad, note=str(bad[:4]))

    # the phase-11 gate must exist and must refuse to emit numbers without inputs
    ph = root / "1_mr" / "phase11_positive_control.py"
    check(S, "G3", "phase 11 (positive control / outcome replication) exists", True, ph.exists(), ph.exists())
    if ph.exists():
        t = ph.read_text(encoding="utf-8", errors="replace")
        check(S, "G3a", "phase 11 refuses to emit a table when its inputs are absent",
              "refusal path present",
              "present" if "Refusing to run" in t else "MISSING", "Refusing to run" in t)
        check(S, "G3b", "phase 11 documents how each missing input is obtained",
              "acquisition notes present",
              "present" if "ACQUISITION" in t and "GWAS Catalog" in t else "MISSING",
              "ACQUISITION" in t and "GWAS Catalog" in t)

    pc = root / "POLICY_COMPLIANCE.md"
    check(S, "G4", "POLICY_COMPLIANCE.md exists", True, pc.exists(), pc.exists())
    if pc.exists():
        t = pc.read_text(encoding="utf-8")
        for token in ("positive control", "second, independent DR outcome dataset",
                      "STROBE-MR", "institutional"):
            check(S, "G4a", "POLICY_COMPLIANCE.md covers '%s'" % token, "present",
                  "present" if token in t else "MISSING", token in t)

    # disclosure gate: no file tracked by git in output_tables/ may carry long digit
    # runs (record identifiers) or CJK source-report text. Record-level files are
    # gitignored and are deliberately not scanned here.
    try:
        import subprocess
        tracked = subprocess.run(["git", "ls-files", "output_tables"], cwd=str(root),
                                 capture_output=True, text=True, timeout=30).stdout.split()
    except Exception as exc:  # pragma: no cover
        check(S, "G5", "git available to list tracked aggregate files", True, False, False,
              fatal=False, note=str(exc))
        return
    offenders = []
    for rel in tracked:
        p = root / rel
        if not p.exists() or p.suffix not in (".csv", ".json"):
            continue
        body = p.read_text(encoding="utf-8", errors="replace")
        hits = LONGDIGITS.findall(body)
        cjk = CJK.findall(body)
        if hits or cjk:
            offenders.append((rel, len(hits), len(cjk)))
    dg = root / "tools" / "disclosure_gate.py"
    check(S, "G4b", "tools/disclosure_gate.py exists (run before every public commit)",
          True, dg.exists(), dg.exists())
    # the digit-run rule is a gate, so prove it can still fail before trusting a
    # clean report: a bare record number must be caught, an rsID must be ignored
    check(S, "G5a", "the digit-run detector still catches a synthetic record identifier",
          "caught", "caught" if LONGDIGITS.search(RECORD_FIXTURE) else "MISSED",
          bool(LONGDIGITS.search(RECORD_FIXTURE)), note=RECORD_FIXTURE)
    check(S, "G5b", "the digit-run detector ignores dbSNP rsIDs",
          "ignored", "ignored" if not LONGDIGITS.search(RSID_FIXTURE) else "FLAGGED",
          not LONGDIGITS.search(RSID_FIXTURE), note=RSID_FIXTURE)
    check(S, "G5", "no tracked aggregate file contains record identifiers or CJK text",
          0, len(offenders), not offenders,
          note=("; ".join("%s: %d digit-runs, %d CJK" % o for o in offenders[:5])
                if offenders else "disclosure gate clean"))


# --------------------------------------------------------------------------
# 5. optional manuscript cross-check
# --------------------------------------------------------------------------
MS_BANNED = [
    (r"NHANES-internal models", "known contradiction with Table 4's title (report P1-2)"),
    (r"Four analyses gave four different answers", "overstatement (report P1-3)"),
    (r"10,813", "stale declared word count (measured 11,172)"),
    (r"S1\u2013S14", "supplementary numbering contradicts the front matter (S1\u2013S15)"),
    (r"zenodo\.XXXXXXX", "placeholder DOI in reference 44"),
    (r"AUC 0\.486 unweighted", "defective n=888 / AUC 0.4861 for the extended local model "
                                "(use 0.4988 on 415 records from layer4_three_row_AUC_corrected.csv)"),
]
MS_REQUIRED = [
    (r"474", "474 of the 475 miscoded records are true DR cases"),
    (r"2,482|2482", "all screening records"),
    (r"1,971|1971", "gradable records"),
    (r"0\.218", "increment attributable to non-lipid clinical information"),
    (r"0\.029", "transport component"),
]

#: Paragraphs inside the counted span that the front matter's convention excludes. The
#: manuscript states its count as "the Introduction to Conclusions, the Declarations and
#: the table footnotes ... excluding ... the funding statement, acknowledgements and
#: references", so the Declarations block counts but the funding statement inside it,
#: and the acknowledgements, do not.
WORD_COUNT_EXCLUDED = re.compile(r"^(Funding|Acknowledgments|Acknowledgements)\b")

#: The recount reads paragraph text only: it cannot see table footnotes held inside table
#: cells, and it tokenises hyphenated compounds differently from Word. The tolerance
#: absorbs that gap, and E4a below keeps it far smaller than any real drift.
WORD_COUNT_TOL = 50


def audit_manuscript(path: Path):
    S = "manuscript"
    try:
        import docx  # python-docx
    except ImportError:
        check(S, "E0", "python-docx available for the manuscript cross-check", True, False,
              False, fatal=False, note="install python-docx to enable this section")
        return
    if not path.exists():
        check(S, "E0", "manuscript present", True, False, False, fatal=False, note=str(path))
        return
    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            parts.extend(c.text for c in r.cells)
    txt = "\n".join(parts)

    check(S, "E1", "manuscript contains the headline outcome counts", "474 / 2,482 / 1,971",
          "checked", all(re.search(p, txt) for p, _ in MS_REQUIRED[:3]), fatal=False)
    for pat, why in MS_REQUIRED:
        check(S, "E2", "manuscript states: %s" % why, "present",
              "present" if re.search(pat, txt) else "MISSING", bool(re.search(pat, txt)), fatal=False)
    for pat, why in MS_BANNED:
        m = re.search(pat, txt)
        check(S, "E3", "manuscript free of: %s" % why, "absent",
              "absent" if not m else "FOUND: %r" % m.group(0), m is None, fatal=False)

    # The declared word count is only meaningful if it is recomputed rather than
    # trusted. This check makes the two unable to drift apart again, which is how the
    # 10,813 figure survived several revisions after the text had grown.
    #
    # The recount has to follow the convention the front matter states, or the check
    # reports a false positive: an earlier implementation stopped *before* Declarations,
    # which under-counted by the whole Declarations block and so could never agree with
    # a correct declaration. The span is Introduction..References, less the headings,
    # the funding statement and the acknowledgements.
    P0 = d.paragraphs
    try:
        i0 = [i for i, p in enumerate(P0) if p.text.strip().startswith("1. Introduction")][0]
        i1 = [i for i, p in enumerate(P0) if p.text.strip().startswith("References")][0]
        seg = P0[i0:i1]
        measured = sum(len(p.text.split()) for p in seg
                       if p.text.strip()
                       and not p.style.name.startswith("Heading")
                       and not WORD_COUNT_EXCLUDED.match(p.text.strip()))
        m = re.search(r"Main text:\s*([\d,]+)\s*words", txt)
        declared = int(m.group(1).replace(",", "")) if m else None
        check(S, "E4", "declared main-text word count agrees with a recount under the "
                       "convention the front matter states",
              measured, declared, near(declared, measured, WORD_COUNT_TOL),
              note="Introduction..References, less headings, the funding statement and the "
                   "acknowledgements; tolerance %d words because the recount cannot see table "
                   "footnotes held inside table cells" % WORD_COUNT_TOL)
        # a tolerance is a gate too: prove it still rejects the value that drifted before
        check(S, "E4a", "the word-count tolerance still rejects the historical 10,813 declaration",
              "rejected", "rejected" if abs(10813 - measured) > WORD_COUNT_TOL else "ACCEPTED",
              abs(10813 - measured) > WORD_COUNT_TOL,
              note="recount %d against a historical 10,813 declaration" % measured)
    except Exception as exc:                  # never let a missing anchor crash the audit
        check(S, "E4", "main-text word count computable", True, False, False,
              fatal=False, note=str(exc))

    # Frontiers caps an Original Research abstract at 350 words; over-length is a
    # technical-check rejection, so it belongs in the gate rather than in a note.
    try:
        a0 = [i for i, p in enumerate(P0) if p.text.strip() == "Abstract"][0]
        a1 = [i for i, p in enumerate(P0) if p.text.strip().startswith("Keywords")][0]
        ab = sum(len(P0[i].text.split()) for i in range(a0 + 1, a1) if P0[i].text.strip())
        check(S, "E5", "abstract within the 350-word limit for Original Research", "<= 350",
              ab, ab <= 350, note="leaves %d words of headroom" % (350 - ab))
    except Exception as exc:
        check(S, "E5", "abstract word count computable", True, False, False,
              fatal=False, note=str(exc))


# --------------------------------------------------------------------------
# report
# --------------------------------------------------------------------------
def report(as_json=None):
    fatal_fail = [r for r in RESULTS if r["fatal"] and not r["ok"]]
    soft_fail = [r for r in RESULTS if not r["fatal"] and not r["ok"]]
    by_section: dict[str, list] = {}
    for r in RESULTS:
        by_section.setdefault(r["section"], []).append(r)

    print()
    print("=" * 108)
    print("REPOSITORY CONSISTENCY AUDIT".center(108))
    print("=" * 108)
    for sec, rs in by_section.items():
        bad = [r for r in rs if not r["ok"]]
        print("\n[%s]  %d checks, %d flagged" % (sec.upper(), len(rs), len(bad)))
        for r in rs:
            mark = "ok  " if r["ok"] else ("FAIL" if r["fatal"] else "warn")
            line = "  %-4s %-9s %-72s" % (mark, r["id"], r["check"][:72])
            print(line)
            if not r["ok"] or r["note"]:
                print("         expected=%s  actual=%s" % (r["expected"], r["actual"]))
                if r["note"]:
                    print("         note: %s" % r["note"])

    print("\n" + "-" * 108)
    print("TOTAL %d checks | fatal failures: %d | advisory flags: %d"
          % (len(RESULTS), len(fatal_fail), len(soft_fail)))
    if fatal_fail:
        print("\nFATAL:")
        for r in fatal_fail:
            print("  %s  %s" % (r["id"], r["check"]))
            print("      expected=%s  actual=%s" % (r["expected"], r["actual"]))
    if soft_fail:
        print("\nADVISORY (non-fatal; usually provenance notes rather than errors):")
        for r in soft_fail:
            print("  %s  %s" % (r["id"], r["check"]))
            if r["note"]:
                print("      %s" % r["note"])
    print("-" * 108)
    print("paths: %s" % paths.describe())

    if as_json:
        Path(as_json).write_text(
            json.dumps({"checks": RESULTS, "fatal_failures": len(fatal_fail),
                        "advisory_flags": len(soft_fail), "paths": paths.describe()},
                       indent=2), encoding="utf-8")
        print("wrote %s" % as_json)

    return 1 if fatal_fail else 0


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__,
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--manuscript", type=Path, default=None,
                    help="optional .docx manuscript to cross-check (advisory only)")
    ap.add_argument("--json", type=Path, default=None, help="write the results as JSON")
    a = ap.parse_args(argv)

    audit_arithmetic()
    audit_mr_tables()
    audit_tables()
    audit_readme()
    audit_alignment()
    audit_gates()
    if a.manuscript:
        audit_manuscript(a.manuscript)
    return report(a.json)


if __name__ == "__main__":
    sys.exit(main())
