# -*- coding: utf-8 -*-
"""Generate Supplementary_Materials.docx (S1-S6) from 方案A_论文图表_Frontiers版.md
   and convert Figure1-4 PNG -> TIFF (300 dpi, LZW) for Frontiers submission."""
import re, os, sys
sys.path.insert(0, r"E:/workbuddy/2026-09-04-18-19-16/figs_frontiers")

TBL = r"C:/Users/Administrator/Desktop/血脂，糖网等数据清洗/清洗结果/方案A_论文图表_Frontiers版.md"
OUT_DOCX = r"E:/workbuddy/Front  in  Endocrinology  投稿/Supplementary_Materials.docx"
SUBDIR = r"E:/workbuddy/Front  in  Endocrinology  投稿"
FIGDIR = r"E:/workbuddy/2026-09-04-18-19-16/figs_frontiers"

# ---------- parse supplementary sections (logic identical to build_docx.py) ----------
def read_lines(p):
    with open(p, encoding="utf-8") as f:
        return f.read().splitlines()

def split_bold(s):
    parts = re.split(r"\*\*(.+?)\*\*", s)
    out = []
    for i, seg in enumerate(parts):
        seg = seg.replace("`", "")
        if seg:
            out.append((seg, i % 2 == 1))
    return out or [(" ", False)]

def parse_md_table_block(lines, i):
    rows = []
    while i < len(lines) and lines[i].strip().startswith("|"):
        cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
            rows.append(cells)
        i += 1
    return rows, i

def parse_table_sections(tbl_md):
    lines = read_lines(tbl_md)
    sections, cur, i = {}, None, 0
    while i < len(lines):
        st = lines[i].strip()
        if st.startswith("## ") and "Table" in st:
            m = re.match(r"##\s+(?:Supplementary\s+)?Table\s+(S?\d)\.", st)
            if m:
                cur = {"title": st[3:].strip(), "elems": []}
                sections[m.group(1)] = cur
                i += 1; continue
        if cur is not None:
            if st.startswith("|"):
                rows, i2 = parse_md_table_block(lines, i)
                if rows:
                    cur["elems"].append(("table", rows))
                i = i2; continue
            if re.match(r"^\*\*\([abc]\)", st):
                cur["elems"].append(("sub", st.strip("*")))
                i += 1; continue
            if st.startswith("*") and st.endswith("*") and not st.startswith("**"):
                cur["elems"].append(("note", st.strip("*").strip()))
                i += 1; continue
            if st.startswith("## ") or st == "---":
                cur = None
        i += 1
    for k, sec in sections.items():
        merged = []
        for e in sec["elems"]:
            if e[0] == "table" and merged and merged[-1][0] == "table":
                merged[-1] = ("table", merged[-1][1] + e[1])
            else:
                merged.append(e)
        sec["elems"] = merged
    return sections

TSEC = parse_table_sections(TBL)
print("parsed sections:", sorted(TSEC.keys()))
for k in sorted(TSEC):
    kinds = [e[0] for e in TSEC[k]["elems"]]
    print(" ", k, kinds)

# ---------- build supplementary docx ----------
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()
stn = doc.styles["Normal"]
stn.font.name = "Times New Roman"; stn.font.size = Pt(10.5)
stn.element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
stn.paragraph_format.space_after = Pt(4)

def para(text="", size=10.5, bold=False, align=None, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    if align is not None:
        p.alignment = align
    for seg, b in split_bold(text):
        r = p.add_run(seg)
        r.font.size = Pt(size); r.font.bold = bold or b
        r.font.name = "Times New Roman"
        r.element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    return p

def set_three_line(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge, sz, val in (("top", 12, "single"), ("bottom", 12, "single"),
                          ("left", 0, "none"), ("right", 0, "none"),
                          ("insideH", 0, "none"), ("insideV", 0, "none")):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), val); el.set(qn("w:sz"), str(sz))
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tblPr.append(borders)
    for cell in table.rows[0].cells:
        tcPr = cell._tc.get_or_add_tcPr()
        tb = OxmlElement("w:tcBorders")
        bot = OxmlElement("w:bottom")
        bot.set(qn("w:val"), "single"); bot.set(qn("w:sz"), "6")
        bot.set(qn("w:color"), "000000")
        tb.append(bot); tcPr.append(tb)

def add_md_table(rows, fontsize=8.5):
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate(rows):
        for ci in range(ncol):
            cell = t.cell(ri, ci)
            pcell = cell.paragraphs[0]
            pcell.paragraph_format.space_after = Pt(1)
            text = row[ci] if ci < len(row) else ""
            for seg, b in split_bold(text):
                r = pcell.add_run(seg)
                r.font.size = Pt(fontsize)
                r.font.bold = (ri == 0) or b
                r.font.name = "Times New Roman"
                r.element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    set_three_line(t)

def add_table_section(key, fontsize=8.5):
    sec = TSEC[key]
    para("**" + sec["title"] + "**", size=10, space_after=6)
    for kind, val in sec["elems"]:
        if kind == "sub":
            para("**" + val + "**", size=8.5, space_after=2)
        elif kind == "table":
            add_md_table(val, fontsize)
            para("", size=4, space_after=2)
        elif kind == "note":
            para("*" + val + "*", size=8.5, space_after=8)

# title page
para("Supplementary Material", size=14, bold=True,
     align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
para("Complete-case analysis inflates discrimination and MR causal signals "
     "are instrument-dependent: a four-layer falsification of lipid-based "
     "diabetic retinopathy risk stratification", size=10.5, bold=True,
     align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
para("Contents: Supplementary Tables S1–S6. Table numbering matches the "
     "in-text citations in the main manuscript.", size=9,
     align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)

for k in ["S1", "S2", "S3", "S4", "S5", "S6"]:
    add_table_section(k)
    para("", size=6, space_after=6)

doc.save(OUT_DOCX)
print("saved:", OUT_DOCX, os.path.getsize(OUT_DOCX), "bytes")

# ---------- PNG -> TIFF (300 dpi, LZW lossless) ----------
from PIL import Image
names = ["Figure1_study_design", "Figure2_MR_forest",
         "Figure3_calibration", "Figure4_ROC"]
IMGDIR = os.path.join(SUBDIR, "手稿1_图片")
os.makedirs(IMGDIR, exist_ok=True)
for n in names:
    src = os.path.join(FIGDIR, n + ".png")
    im = Image.open(src)
    if im.mode != "RGB":                      # journals: no alpha channel
        im = im.convert("RGB")
    m = re.match(r"Figure(\d)", n)
    for dst in (os.path.join(SUBDIR, f"Figure{m.group(1)}.tiff"),
                os.path.join(IMGDIR, f"Figure{m.group(1)}.tiff")):
        im.save(dst, format="TIFF", compression="tiff_lzw", dpi=(300, 300))
        print(f"tiff: {dst} ({os.path.getsize(dst)//1024} KB, {im.size[0]}x{im.size[1]}px, dpi=300, RGB)")
